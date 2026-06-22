import os
import io
import calendar
import re
from functools import wraps
from collections import defaultdict
from datetime import datetime, timedelta
import uuid

from flask import Blueprint, jsonify, render_template, request, send_file

from database import (
    attach_job_order_transaction,
    claim_job_order_for_activation,
    clear_machine_run_window,
    create_post_cycle_payment_log,
    create_job_order,
    create_job_orders_bulk,
    delete_job_order,
    delete_transactions,
    delete_machines,
    get_active_shift,
    get_all_machines,
    get_active_service_by_name,
    get_connection,
    get_email_settings,
    get_dashboard_layout,
    get_job_order,
    get_machine,
    get_receipt_format_config,
    get_recent_shift_summary_count,
    get_recent_transactions,
    get_shift,
    get_services_by_ids,
    get_transaction_by_request_id,
    increment_customer_order_count,
    insert_transaction_with_items,
    list_open_job_orders,
    summarize_post_cycle_payment_logs,
    list_customers,
    list_manual_expenses,
    set_email_settings,
    set_machine_run_window,
    upsert_customer,
    set_recent_shift_summary_count,
    set_receipt_format_config,
    set_dashboard_layout,
    update_machine,
    update_transaction_gcash_amount,
    save_receipt_overrides,
    revert_job_order_usage,
    update_machine_status,
    list_promos,
    create_promo,
    update_promo,
    delete_promo,
    get_day_change_time,
    set_day_change_time,
    get_analytics_settings,
    set_analytics_settings,
    get_time_shifts,
    list_post_cycle_payment_logs,
)

from routes.security import require_admin_pin
from services.esp32 import check_esp32_life, get_esp32_status, send_pulse, async_send_pulse
from services.sync import try_immediate_sync


dashboard_bp = Blueprint("dashboard", __name__)

IS_DEV = os.environ.get("FLASK_ENV", "development") == "development"
DEFAULT_LOCATION_ID = os.environ.get("LOCATION_ID", "local")
MACHINE_RUN_SECONDS = int(os.environ.get("MACHINE_RUN_SECONDS", str(35 * 60)))
EXTRA_WASH_NAME = "extra wash"
EXTRA_DRY_NAME = "extra dry"
QUICK_SERVICE_NAMES = {EXTRA_WASH_NAME, EXTRA_DRY_NAME}
PHONE_PATTERN = re.compile(r"^\d{10,11}$")


def _now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _operational_day_window(day_str):
    change_time = get_day_change_time()
    try:
        parts = change_time.split(":")
        hh = int(parts[0])
        mm = int(parts[1])
    except Exception:
        hh = 8
        mm = 1

    dt_start = datetime.strptime(day_str, "%Y-%m-%d") + timedelta(hours=hh, minutes=mm)
    dt_end = dt_start + timedelta(days=1) - timedelta(seconds=1)
    return (
        dt_start.strftime("%Y-%m-%d %H:%M:%S"),
        dt_end.strftime("%Y-%m-%d %H:%M:%S")
    )


def _current_operational_day_str():
    change_time = get_day_change_time()
    try:
        parts = change_time.split(":")
        hh = int(parts[0])
        mm = int(parts[1])
    except Exception:
        hh = 8
        mm = 1
    now = datetime.now()
    day_start_today = now.replace(hour=hh, minute=mm, second=0, microsecond=0)
    if now < day_start_today:
        return (now - timedelta(days=1)).strftime("%Y-%m-%d")
    return now.strftime("%Y-%m-%d")


def _json_api_guard(handler):
    @wraps(handler)
    def wrapped(*args, **kwargs):
        try:
            return handler(*args, **kwargs)
        except Exception as exc:
            return jsonify({"error": f"Internal server error: {exc}"}), 500

    return wrapped


def _is_valid_ipv4(value):
    parts = str(value or "").strip().split(".")
    if len(parts) != 4:
        return False
    for part in parts:
        if not part.isdigit():
            return False
        number = int(part)
        if number < 0 or number > 255:
            return False
    return True


def _exclude_quick_service_items(sale_items):
    service_items = [
        item for item in (sale_items or [])
        if str(item.get("kind") or "").strip().lower() == "service"
    ]
    if not service_items:
        return sale_items or []

    service_ids = [str(item.get("item_id") or "").strip() for item in service_items]
    service_map = get_services_by_ids(service_ids)
    filtered = []

    for item in sale_items or []:
        if str(item.get("kind") or "").strip().lower() != "service":
            filtered.append(item)
            continue

        service_id = str(item.get("item_id") or "").strip()
        service = service_map.get(service_id)
        service_name = str((service or {}).get("name") or "").strip().lower()
        if service_name in QUICK_SERVICE_NAMES:
            continue
        filtered.append(item)

    return filtered


def _quick_service_name_for_machine(machine_type):
    mtype = str(machine_type or "").strip().lower()
    if mtype == "washer":
        return "Extra Wash"
    if mtype == "dryer":
        return "Extra Dry"
    return None


def _resolve_customer_payload(data):
    payload = data if isinstance(data, dict) else {}
    customer = payload.get("customer") if isinstance(payload.get("customer"), dict) else {}

    name = str(customer.get("name") or "").strip()
    phone_raw = customer.get("phone")
    phone = str(phone_raw).strip() if phone_raw not in (None, "") else None

    if not name:
        raise ValueError("Customer name is required")
    if phone and not PHONE_PATTERN.fullmatch(phone):
        raise ValueError("Customer phone must be 10 to 11 digits")

    return upsert_customer(name=name, phone=phone)


def _job_order_error_status(error_message):
    msg = str(error_message or "").strip().lower()
    if "already used" in msg or "closed" in msg:
        return 409
    return 400


def _price_by_machine_type(machine_type, mode_id, machine=None):
    mtype = str(machine_type or "").strip().lower()
    mode_id = str(mode_id or "standard").strip().lower()

    if not machine:
        machines = get_all_machines()
        matching = [m for m in machines if str(m.get("type") or "").strip().lower() == mtype]
        if matching:
            machine = matching[0]

    if not machine:
        return 0

    custom_modes_str = machine.get("custom_modes")
    if custom_modes_str:
        try:
            import json
            modes = json.loads(custom_modes_str)
            for m in modes:
                if str(m.get("id")).lower() == mode_id:
                    return max(0, int(m.get("price") or 0))
        except Exception:
            pass
            
    if mode_id == "quick":
        return max(0, int(machine.get("quick_wash_price") or machine.get("vend_price") or 0))
    return max(0, int(machine.get("vend_price") or 0))


def _normalize_job_order_mode(mode_value):
    return str(mode_value or "standard").strip().lower()


def _resolve_activation_profile(machine, activation_mode):
    selected_mode = str(activation_mode or "standard").strip().lower()
    custom_modes_json = machine.get("custom_modes")
    
    if custom_modes_json:
        try:
            import json
            modes = json.loads(custom_modes_json)
            for m in modes:
                if str(m.get("id")).lower() == selected_mode:
                    pulse_count = max(1, int(m.get("pulse_count") or 1))
                    base_amount = max(0, int(m.get("price") or 0))
                    return selected_mode, pulse_count, base_amount, m.get("products") or [], m.get("services") or []
        except Exception:
            pass

    # Fallback to legacy
    if selected_mode not in ("standard", "quick"):
        selected_mode = "standard"

    if selected_mode == "quick":
        pulse_count = max(1, int(machine.get("quick_wash_pulse_count") or 1))
        base_amount = max(0, int(machine.get("quick_wash_price") or machine.get("vend_price") or 0))
    else:
        pulse_count = int(machine.get("pulse_count") or 2)
        base_amount = max(0, int(machine.get("vend_price") or 0))

    return selected_mode, pulse_count, base_amount, [], []


def _activate_machine_with_sale(
    machine,
    active_shift,
    sale_items,
    request_id,
    customer_payload,
    paid_by_gcash,
    activation_mode,
):
    selected_mode, pulse_count, base_amount, custom_products, custom_services = _resolve_activation_profile(machine, activation_mode)
    
    # Merge bundled items from custom modes
    for cp in custom_products:
        item_id = str(cp.get("id") or "").strip()
        qty = int(cp.get("qty") or 0)
        if item_id and qty > 0:
            sale_items.append({
                "kind": "product",
                "item_id": item_id,
                "quantity": qty,
                "unit_price": 0
            })
    for cs in custom_services:
        item_id = str(cs.get("id") or "").strip()
        qty = int(cs.get("qty") or 0)
        if item_id and qty > 0:
            sale_items.append({
                "kind": "service",
                "item_id": item_id,
                "quantity": qty,
                "unit_price": 0
            })

    timestamp = _now_str()
    import uuid

    txn_id = str(uuid.uuid4())
    txn_status = "COMPLETED"
    txn_result = insert_transaction_with_items(
        txn_id,
        machine["id"],
        base_amount,
        txn_status,
        timestamp,
        employee_id=active_shift["employee_id"],
        shift_id=active_shift["id"],
        sale_items=sale_items,
        request_id=request_id,
        customer_id=customer_payload.get("customer_id") if customer_payload else None,
        customer_name=customer_payload.get("name") if customer_payload else None,
        customer_phone=customer_payload.get("phone") if customer_payload else None,
        paid_by_gcash=paid_by_gcash,
    )

    if customer_payload and customer_payload.get("customer_id"):
        increment_customer_order_count(
            customer_id=customer_payload["customer_id"],
            machine_type=machine.get("type"),
            quantity=1
        )

    run_ends_at = (datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S") + timedelta(seconds=MACHINE_RUN_SECONDS)).strftime("%Y-%m-%d %H:%M:%S")
    set_machine_run_window(machine["id"], timestamp, run_ends_at)

    final_txn_id = txn_result.get("transaction_id", txn_id)

    def on_result(success, message):
        if not success and not IS_DEV:
            delete_transactions([final_txn_id])
            clear_machine_run_window(machine["id"])
            update_machine_status(machine["id"], "OFFLINE")

    async_send_pulse(
        machine["esp32_ip"],
        machine["pulse_on"],
        machine["pulse_off"],
        pulse_count,
        on_result
    )

    return {
        "status": txn_status,
        "transaction_id": final_txn_id,
        "stored_transaction_id": final_txn_id,
        "machine": machine["name"],
        "machine_id": machine["id"],
        "amount": txn_result["total_amount"],
        "base_amount": base_amount,
        "activation_mode": selected_mode,
        "product_total": txn_result["product_total"],
        "service_total": txn_result["service_total"],
        "item_count": txn_result["item_count"],
        "low_stock_warnings": txn_result["low_stock_warnings"],
        "idempotent_hit": txn_result.get("idempotent_hit", False),
        "remaining_seconds": MACHINE_RUN_SECONDS,
        "employee_id": active_shift["employee_id"],
        "employee_name": active_shift.get("display_name"),
        "shift_id": active_shift["id"],
        "customer": customer_payload,
        "customer_id": customer_payload.get("customer_id") if customer_payload else None,
        "customer_name": customer_payload.get("name") if customer_payload else None,
        "customer_phone": customer_payload.get("phone") if customer_payload else None,
        "paid_by_gcash": int(txn_result.get("paid_by_gcash") or 0),
        "job_order_id": None,
        "job_order_no": None,
        "job_order_status": None,
        "job_order_remaining_wash_qty": 0,
        "job_order_remaining_dry_qty": 0,
    }


def _get_transactions(limit=None, offset=0):
    return get_recent_transactions(limit=limit, offset=offset)


def _enrich_machines(raw_machines):
    enriched = []
    for m in raw_machines:
        row = dict(m)
        row["location_id"] = DEFAULT_LOCATION_ID
        row["location_name"] = "Local Pi"
        enriched.append(row)
    return enriched


def _enrich_transactions(raw_transactions):
    enriched = []
    all_txn_ids = []

    for t in raw_transactions:
        row = dict(t)
        request_id = str(row.get("request_id") or "").strip()
        is_bulk = request_id.startswith("bulk-")
        bulk_group_id = ""
        if is_bulk:
            bulk_group_id = request_id.split(":", 1)[0]
        row["location_id"] = DEFAULT_LOCATION_ID
        row["location_name"] = "Local Pi"
        row["is_bulk_activation"] = is_bulk
        row["bulk_group_id"] = bulk_group_id
        row["addon_items"] = []
        tid = str(row.get("id") or "").strip()
        if tid:
            all_txn_ids.append(tid)
        enriched.append(row)

    # Batch-fetch items for all transactions in one query
    if all_txn_ids:
        conn = get_connection()
        try:
            placeholders = ",".join("?" for _ in all_txn_ids)
            rows = conn.execute(
                f"""
                SELECT transaction_id, item_type, item_name, quantity, unit_price, line_total
                FROM transaction_items
                WHERE transaction_id IN ({placeholders})
                ORDER BY transaction_id, item_type, item_name
                """,
                all_txn_ids,
            ).fetchall()
        finally:
            conn.close()

        items_by_txn = {}
        for r in rows:
            tid = r["transaction_id"]
            if tid not in items_by_txn:
                items_by_txn[tid] = []
            items_by_txn[tid].append({
                "item_type": r["item_type"],
                "item_name": r["item_name"],
                "quantity": int(r["quantity"] or 0),
                "unit_price": int(r["unit_price"] or 0),
                "line_total": int(r["line_total"] or 0),
            })

        for row in enriched:
            tid = str(row.get("id") or "").strip()
            if tid in items_by_txn:
                row["addon_items"] = items_by_txn[tid]

    return enriched


def _build_dashboard_stats(transactions, machines):
    today = datetime.now().date()
    today_revenue = 0

    for t in transactions:
        started = t.get("started_at")
        if started:
            started = started[:10]

        if started == str(today) and t.get("status") == "COMPLETED":
            today_revenue += int(t.get("amount") or 0)

    return {
        "today_revenue": today_revenue,
        "total_transactions": len(transactions),
        "location_count": 1,
        "machine_count": len(machines),
    }


def _filter_by_date(transactions, start_date=None, end_date=None):
    if not start_date and not end_date:
        return transactions

    filtered = []
    for t in transactions:
        try:
            ts = datetime.strptime(t["started_at"], "%Y-%m-%d %H:%M:%S").date()
        except (ValueError, TypeError):
            continue
        if start_date and ts < start_date:
            continue
        if end_date and ts > end_date:
            continue
        filtered.append(t)
    return filtered


def _sum_manual_expenses(start_date=None, end_date=None):
    conn = get_connection()
    try:
        if start_date or end_date:
            from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
            to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
            row = conn.execute(
                "SELECT SUM(amount) AS total FROM manual_expenses WHERE expense_at >= ? AND expense_at <= ?",
                (from_str, to_str),
            ).fetchone()
        else:
            row = conn.execute("SELECT SUM(amount) AS total FROM manual_expenses").fetchone()
        return int((row["total"] if row else 0) or 0)
    finally:
        conn.close()


def _group_manual_expenses(start_date=None, end_date=None):
    conn = get_connection()
    try:
        if start_date or end_date:
            from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
            to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
            rows = conn.execute(
                """
                SELECT COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense') AS name,
                       COUNT(*) AS count,
                       SUM(amount) AS total
                FROM manual_expenses
                WHERE expense_at >= ? AND expense_at <= ?
                GROUP BY COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense')
                ORDER BY name ASC
                """,
                (from_str, to_str),
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense') AS name,
                       COUNT(*) AS count,
                       SUM(amount) AS total
                FROM manual_expenses
                GROUP BY COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense')
                ORDER BY name ASC
                """
            ).fetchall()
    finally:
        conn.close()

    return [
        {
            "name": row["name"],
            "count": int(row["count"] or 0),
            "total": int(row["total"] or 0),
        }
        for row in rows
    ]


def _build_machine_usage_breakdown(transactions, machines):
    machine_meta = {
        str(m.get("id") or "").strip(): {
            "name": m.get("name") or str(m.get("id") or "Unknown"),
            "type": str(m.get("type") or "").strip().lower() or "unknown",
        }
        for m in machines or []
    }

    usage = {}
    for txn in transactions or []:
        machine_id = str(txn.get("machine_id") or "").strip()
        if not machine_id or machine_id == "post-cycle-addons":
            continue

        amount = int(txn.get("amount") or 0)
        product_total = int(txn.get("product_total") or 0)
        service_total = int(txn.get("service_total") or 0)
        machine_amount = max(0, amount - product_total - service_total)

        if machine_id not in usage:
            meta = machine_meta.get(machine_id, {"name": machine_id, "type": "unknown"})
            usage[machine_id] = {
                "machine_id": machine_id,
                "machine_name": meta.get("name") or machine_id,
                "machine_type": meta.get("type") or "unknown",
                "count": 0,
                "revenue": 0,
            }

        usage[machine_id]["count"] += 1
        usage[machine_id]["revenue"] += machine_amount

    rows = list(usage.values())
    rows.sort(key=lambda row: row["count"], reverse=True)
    return rows


def _build_analytics_receipt_snapshot(transactions, machines, start_date=None, end_date=None):
    valid_sales = [
        t for t in transactions
        if t.get("status") in ("COMPLETED", "SIMULATED")
    ]

    txn_ids = [str(t.get("id") or "").strip() for t in valid_sales if str(t.get("id") or "").strip()]

    # Fetch job orders first so we can extract their IDs
    if start_date or end_date:
        from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
        to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
        job_orders = _fetch_job_orders_for_day(from_str, to_str)
    else:
        job_orders = _fetch_job_orders_for_day("0000-01-01 00:00:00", "9999-12-31 23:59:59")
        
    job_order_ids = [jo["id"] for jo in job_orders if jo.get("id")]

    item_usage = _fetch_item_usage(txn_ids)
    item_breakdown = _fetch_item_breakdown(txn_ids)

    gross_collected = sum(int(t.get("amount") or 0) for t in valid_sales)
    product_revenue = sum(int(t.get("product_total") or 0) for t in valid_sales)
    # Track transaction-native service revenue separately — these are already embedded in
    # gross_collected (txn.amount includes them), so they're the only portion valid for
    # cash/net adjustments.
    txn_service_revenue = sum(int(t.get("service_total") or 0) for t in valid_sales)
    service_revenue = txn_service_revenue

    txn_gcash_revenue = 0
    txn_gcash_count = 0
    for txn in valid_sales:
        amount = int(txn.get("amount") or 0)
        if txn.get("gcash_amount") is not None:
            gcash_amount = int(txn.get("gcash_amount") or 0)
        else:
            gcash_amount = amount if int(txn.get("paid_by_gcash") or 0) == 1 else 0
        gcash_amount = max(0, min(gcash_amount, amount))
        txn_gcash_revenue += gcash_amount
        if gcash_amount > 0:
            txn_gcash_count += 1

    if start_date or end_date:
        from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
        to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
        post_cycle = summarize_post_cycle_payment_logs(start_at=from_str, end_at=to_str)
        report_window = f"{from_str} to {to_str}"
    else:
        post_cycle = summarize_post_cycle_payment_logs()
        report_window = "ALL"

    post_cycle_transfer_amount = int(post_cycle.get("amount") or 0)
    post_cycle_transfer_count = int(post_cycle.get("count") or 0)

    total_sales = gross_collected
    machine_revenue = max(0, total_sales - product_revenue - service_revenue)

    manual_expenses = _sum_manual_expenses(start_date=start_date, end_date=end_date)
    total_expenses = manual_expenses + int(item_usage.get("cogs_total") or 0)
    gcash_revenue = txn_gcash_revenue + post_cycle_transfer_amount

    analytics_settings = get_analytics_settings()
    include_service_in_net = bool(analytics_settings.get("include_service_revenue_in_net", False))
    _svc_adj = 0 if include_service_in_net else -txn_service_revenue

    cash_sales = total_sales + _svc_adj - gcash_revenue
    cash_revenue = cash_sales - total_expenses
    net_sales = total_sales + _svc_adj - total_expenses

    expense_breakdown = _group_manual_expenses(start_date=start_date, end_date=end_date)
    if int(item_usage.get("cogs_total") or 0) > 0:
        expense_breakdown.append(
            {
                "name": "COGS",
                "count": int(item_usage.get("product_qty") or 0),
                "total": int(item_usage.get("cogs_total") or 0),
            }
        )

    return {
        "report_window": report_window,
        "transaction_count": len(valid_sales),
        "gross_collected": gross_collected,
        "total_sales": total_sales,
        "machine_revenue": machine_revenue,
        "cash_sales": cash_sales,
        "cash_revenue": cash_revenue,
        "gcash_job_order_count": txn_gcash_count + post_cycle_transfer_count,
        "gcash_revenue": gcash_revenue,
        "product_revenue": product_revenue,
        "service_revenue": service_revenue,
        "manual_expenses": manual_expenses,
        "cogs_total": int(item_usage.get("cogs_total") or 0),
        "total_expenses": total_expenses,
        "net_sales": total_sales - total_expenses,
        "post_cycle_transfer_amount": post_cycle_transfer_amount,
        "post_cycle_transfer_count": post_cycle_transfer_count,
        "machine_usage_breakdown": _build_machine_usage_breakdown(valid_sales, machines),
        "product_breakdown": item_breakdown.get("product") or [],
        "service_breakdown": item_breakdown.get("service") or [],
        "expense_breakdown": expense_breakdown,
    }


def _is_time_in_shift(dt_str, start_time_str, end_time_str):
    if not dt_str:
        return False
    try:
        if " " in dt_str:
            time_part = dt_str.split(" ")[1]
        else:
            time_part = dt_str
        parts = time_part.split(":")
        t_val = int(parts[0]) * 60 + int(parts[1])
        
        sparts = start_time_str.split(":")
        s_val = int(sparts[0]) * 60 + int(sparts[1])
        
        eparts = end_time_str.split(":")
        e_val = int(eparts[0]) * 60 + int(eparts[1])
        
        if s_val < e_val:
            return s_val <= t_val < e_val
        elif s_val > e_val:
            return t_val >= s_val or t_val < e_val
        else:
            return True
    except Exception:
        return False


def _get_shift_for_time(dt_str, time_shifts):
    for shift in time_shifts:
        if _is_time_in_shift(dt_str, shift["start_time"], shift["end_time"]):
            return shift
    return None


def _get_manual_expenses_list(start_date=None, end_date=None):
    conn = get_connection()
    try:
        if start_date or end_date:
            from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
            to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
            rows = conn.execute(
                "SELECT * FROM manual_expenses WHERE expense_at >= ? AND expense_at <= ? ORDER BY expense_at DESC",
                (from_str, to_str),
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM manual_expenses ORDER BY expense_at DESC").fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def _build_analytics(transactions, machines, start_date=None, end_date=None):
    valid_sales = [t for t in transactions if t.get("status") == "COMPLETED"]
    total_revenue = sum(int(t.get("amount") or 0) for t in valid_sales)
    total_cycles = len(transactions)
    avg_per_cycle = round(total_revenue / total_cycles) if total_cycles else 0
    gcash_collected = 0
    gcash_transaction_count = 0

    revenue_by_day_map = defaultdict(int)
    cycles_by_day_map = defaultdict(int)
    usage_map = defaultdict(lambda: {"cycles": 0, "revenue": 0})

    machine_name_by_id = {m["id"]: m.get("name", m["id"]) for m in machines}

    for t in valid_sales:
        eff_date_str = t.get("started_at")
        machine_id = t.get("machine_id")
        amount = int(t.get("amount") or 0)
        status = t.get("status")

        try:
            day = datetime.strptime(eff_date_str, "%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d")
        except (ValueError, TypeError):
            continue

        cycles_by_day_map[day] += 1
        usage_map[machine_id]["cycles"] += 1

        if status == "COMPLETED":
            revenue_by_day_map[day] += amount
            usage_map[machine_id]["revenue"] += amount
            if t.get("gcash_amount") is not None:
                gcash = int(t.get("gcash_amount") or 0)
            else:
                gcash = amount if int(t.get("paid_by_gcash") or 0) == 1 else 0
            bounded_gcash = max(0, min(gcash, amount))
            gcash_collected += bounded_gcash
            if bounded_gcash > 0:
                gcash_transaction_count += 1
                    
    if start_date or end_date:
        from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
        to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
        post_cycle = summarize_post_cycle_payment_logs(start_at=from_str, end_at=to_str)
    else:
        post_cycle = summarize_post_cycle_payment_logs()

    post_cycle_amount = int(post_cycle.get("amount") or 0)
    post_cycle_count = int(post_cycle.get("count") or 0)
    gcash_collected += post_cycle_amount
    gcash_transaction_count += post_cycle_count

    # NOTE: The linked transaction already has gcash_amount / paid_by_gcash set correctly,
    # so we do NOT add jo.total_amount here — that would double-count GCash revenue.

    receipt_data = _build_analytics_receipt_snapshot(
        transactions,
        machines,
        start_date=start_date,
        end_date=end_date,
    )

    total_expenses = receipt_data["total_expenses"]
    cash_sales = receipt_data["cash_sales"]
    cash_revenue = receipt_data["cash_revenue"]

    revenue_by_day = [
        {"day": day, "revenue": revenue_by_day_map[day]}
        for day in sorted(revenue_by_day_map.keys())
    ]
    cycles_by_day = [
        {"day": day, "cycles": cycles_by_day_map[day]}
        for day in sorted(cycles_by_day_map.keys())
    ]

    machine_usage = []
    for machine_id, vals in usage_map.items():
        machine_usage.append({
            "machine_id": machine_id,
            "name": machine_name_by_id.get(machine_id, machine_id),
            "cycles": vals["cycles"],
            "revenue": vals["revenue"],
        })
    machine_usage.sort(key=lambda x: x["cycles"], reverse=True)

    gcash_share_pct = round((min(gcash_collected, total_revenue) / total_revenue) * 100) if total_revenue else 0

    # Calculate time shift analytics
    time_shifts = get_time_shifts()
    shift_buckets = {}
    for ts in time_shifts:
        shift_buckets[ts["id"]] = {
            "id": ts["id"],
            "name": ts["name"],
            "start_time": ts["start_time"],
            "end_time": ts["end_time"],
            "transaction_count": 0,
            "gross_collected": 0,
            "cash_sales": 0,
            "gcash_revenue": 0,
            "product_revenue": 0,
            "service_revenue": 0,
            "machine_revenue": 0,
            "manual_expenses": 0,
            "cogs_total": 0,
            "total_expenses": 0,
            "net_sales": 0,
            "post_cycle_transfer_amount": 0,
            "post_cycle_transfer_count": 0,
            "txn_ids": []
        }

    unassigned_bucket = {
        "id": None,
        "name": "Other / Unassigned",
        "start_time": "",
        "end_time": "",
        "transaction_count": 0,
        "gross_collected": 0,
        "cash_sales": 0,
        "gcash_revenue": 0,
        "product_revenue": 0,
        "service_revenue": 0,
        "machine_revenue": 0,
        "manual_expenses": 0,
        "cogs_total": 0,
        "total_expenses": 0,
        "net_sales": 0,
        "post_cycle_transfer_amount": 0,
        "post_cycle_transfer_count": 0,
        "txn_ids": []
    }

    # 1. Partition completed transactions
    for t in valid_sales:
        start_time_str = t.get("started_at")
        matched_shift = _get_shift_for_time(start_time_str, time_shifts)
        bucket = shift_buckets[matched_shift["id"]] if matched_shift else unassigned_bucket

        amount = int(t.get("amount") or 0)
        bucket["transaction_count"] += 1
        bucket["gross_collected"] += amount
        bucket["product_revenue"] += int(t.get("product_total") or 0)
        bucket["service_revenue"] += int(t.get("service_total") or 0)
        bucket["txn_ids"].append(t.get("id"))

        if t.get("gcash_amount") is not None:
            gcash_amount = int(t.get("gcash_amount") or 0)
        else:
            gcash_amount = amount if int(t.get("paid_by_gcash") or 0) == 1 else 0
        gcash_amount = max(0, min(gcash_amount, amount))
        bucket["gcash_revenue"] += gcash_amount

    # 2. Partition post-cycle payment logs
    if start_date or end_date:
        from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
        to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
        post_cycle_logs = list_post_cycle_payment_logs(start_at=from_str, end_at=to_str)
    else:
        post_cycle_logs = list_post_cycle_payment_logs()

    for log in post_cycle_logs:
        log_time_str = log.get("logged_at")
        matched_shift = _get_shift_for_time(log_time_str, time_shifts)
        bucket = shift_buckets[matched_shift["id"]] if matched_shift else unassigned_bucket

        log_amount = int(log.get("amount") or 0)
        bucket["post_cycle_transfer_amount"] += log_amount
        bucket["post_cycle_transfer_count"] += 1

    # 3. Partition manual expenses
    expenses_list = _get_manual_expenses_list(start_date, end_date)
    for exp in expenses_list:
        exp_time_str = exp.get("expense_at")
        matched_shift = _get_shift_for_time(exp_time_str, time_shifts)
        bucket = shift_buckets[matched_shift["id"]] if matched_shift else unassigned_bucket

        exp_amount = int(exp.get("amount") or 0)
        bucket["manual_expenses"] += exp_amount

    # 4. Fetch COGS and calculate derived metrics
    analytics_settings = get_analytics_settings()
    include_service_in_net = bool(analytics_settings.get("include_service_revenue_in_net", False))

    all_buckets = list(shift_buckets.values())
    if (unassigned_bucket["transaction_count"] > 0
        or unassigned_bucket["manual_expenses"] > 0
        or unassigned_bucket["post_cycle_transfer_count"] > 0):
        all_buckets.append(unassigned_bucket)

    for bucket in all_buckets:
        if bucket["txn_ids"]:
            item_usage = _fetch_item_usage(bucket["txn_ids"])
            bucket["cogs_total"] = int(item_usage.get("cogs_total") or 0)

        bucket["machine_revenue"] = max(0, bucket["gross_collected"] - bucket["product_revenue"] - bucket["service_revenue"])
        total_gcash = bucket["gcash_revenue"] + bucket["post_cycle_transfer_amount"]
        bucket["gcash_revenue"] = total_gcash

        svc_adj = 0 if include_service_in_net else -bucket["service_revenue"]
        bucket["cash_sales"] = bucket["gross_collected"] + svc_adj - total_gcash
        bucket["total_expenses"] = bucket["manual_expenses"] + bucket["cogs_total"]
        bucket["net_sales"] = bucket["gross_collected"] + svc_adj - bucket["total_expenses"]
        
        # Clean up txn_ids to keep JSON small
        del bucket["txn_ids"]

    return {
        "total_revenue": total_revenue,
        "total_cycles": total_cycles,
        "avg_per_cycle": avg_per_cycle,
        "cash_sales": cash_sales,
        "cash_revenue": cash_revenue,
        "gcash_collected": gcash_collected,
        "gcash_transaction_count": gcash_transaction_count,
        "total_expenses": total_expenses,
        "post_cycle_transfer_amount": post_cycle_amount,
        "post_cycle_transfer_count": post_cycle_count,
        "gcash_share_pct": gcash_share_pct,
        "revenue_by_day": revenue_by_day,
        "cycles_by_day": cycles_by_day,
        "machine_usage": machine_usage,
        "receipt_data": receipt_data,
        "time_shift_analytics": all_buckets,
    }



def _normalize_machine_runtime(machine):
    row = dict(machine)
    status = row.get("status")
    run_ends_at = row.get("run_ends_at")
    if status != "BUSY" or not run_ends_at:
        row["remaining_seconds"] = 0
        return row

    try:
        end_dt = datetime.strptime(run_ends_at, "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        clear_machine_run_window(row["id"])
        row["status"] = "IDLE"
        row["run_started_at"] = None
        row["run_ends_at"] = None
        row["remaining_seconds"] = 0
        return row

    remaining = int((end_dt - datetime.now()).total_seconds())
    if remaining <= 0:
        clear_machine_run_window(row["id"])
        row["status"] = "IDLE"
        row["run_started_at"] = None
        row["run_ends_at"] = None
        row["remaining_seconds"] = 0
        return row

    row["remaining_seconds"] = remaining
    return row


def _fetch_transactions_by_day(day_str):
    day_start, day_end = _operational_day_window(day_str)
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT *
        FROM transactions
        WHERE started_at >= ? AND started_at <= ?
          AND status IN ('COMPLETED', 'SIMULATED')
        ORDER BY started_at ASC
        """,
        (day_start, day_end),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def _get_transactions_for_date_range(start_date=None, end_date=None):
    if not start_date and not end_date:
        return _get_transactions()

    from_str = f"{start_date.strftime('%Y-%m-%d')} 00:00:00" if start_date else "0000-01-01 00:00:00"
    to_str = f"{end_date.strftime('%Y-%m-%d')} 23:59:59" if end_date else "9999-12-31 23:59:59"
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT *
        FROM transactions
        WHERE started_at >= ? AND started_at <= ?
          AND status IN ('COMPLETED', 'SIMULATED')
        ORDER BY started_at ASC
        """,
        (from_str, to_str),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _fetch_transactions_by_shift(shift_id):
    conn = get_connection()
    rows = conn.execute(
        "SELECT * FROM transactions WHERE shift_id = ? ORDER BY started_at ASC",
        (shift_id,),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _fetch_item_usage(txn_ids):
    if not txn_ids:
        return {
            "product_qty": 0,
            "service_qty": 0,
            "product_lines": 0,
            "service_lines": 0,
            "cogs_total": 0,
        }

    conn = get_connection()
    placeholders = ",".join("?" for _ in txn_ids)
    rows = conn.execute(
        f"""
        SELECT item_type, quantity, line_cost
        FROM transaction_items
        WHERE transaction_id IN ({placeholders})
        """,
        txn_ids,
    ).fetchall()
    conn.close()

    product_qty = 0
    service_qty = 0
    product_lines = 0
    service_lines = 0
    cogs_total = 0

    for row in rows:
        kind = row["item_type"]
        qty = int(row["quantity"] or 0)
        if kind == "product":
            product_qty += qty
            product_lines += 1
            cogs_total += int(row["line_cost"] or 0)
        elif kind == "service":
            service_qty += qty
            service_lines += 1

    return {
        "product_qty": product_qty,
        "service_qty": service_qty,
        "product_lines": product_lines,
        "service_lines": service_lines,
        "cogs_total": cogs_total,
    }


def _fetch_item_breakdown(txn_ids):
    if not txn_ids:
        return {"product": [], "service": []}

    conn = get_connection()
    placeholders = ",".join("?" for _ in txn_ids)
    raw_rows = conn.execute(
        f"""
        SELECT item_type, item_name, quantity AS qty, line_total AS total
        FROM transaction_items
        WHERE transaction_id IN ({placeholders})
        """,
        txn_ids,
    ).fetchall()
    conn.close()

    aggregated = {}
    for row in raw_rows:
        key = (row["item_type"], row["item_name"])
        if key not in aggregated:
            aggregated[key] = {"qty": 0, "total": 0}
        aggregated[key]["qty"] += int(row["qty"] or 0)
        aggregated[key]["total"] += int(row["total"] or 0)

    out = {"product": [], "service": []}
    for (kind, name), data in aggregated.items():
        if kind not in out:
            continue
        out[kind].append({
            "name": name,
            "qty": data["qty"],
            "total": data["total"],
        })
        
    out["product"].sort(key=lambda x: x["name"])
    out["service"].sort(key=lambda x: x["name"])
    
    return out


def _compute_cash_sales_and_net(total_sales, txn_gcash_revenue, post_cycle_transfer_amount, total_expenses, service_adjustment=0):
    effective_sales = int(total_sales or 0) + int(service_adjustment or 0)
    cash_sales = effective_sales - int(txn_gcash_revenue or 0) - int(post_cycle_transfer_amount or 0)
    net_cash = cash_sales - int(total_expenses or 0)
    return cash_sales, net_cash


def _fetch_job_orders_for_shift(shift):
    if not shift or not shift.get("id"):
        return []

    end_at = shift.get("ended_at") or _now_str()
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT *
        FROM job_orders
        WHERE created_by_shift_id = ?
           OR (
                created_by_shift_id IS NULL
                AND created_by_employee_id = ?
                AND created_at >= ?
                AND created_at <= ?
           )
        ORDER BY created_at ASC
        """,
        (shift["id"], shift.get("employee_id"), shift.get("started_at"), end_at),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _fetch_job_orders_for_day(day_start, day_end):
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT *
        FROM job_orders
        WHERE created_at >= ? AND created_at <= ?
        ORDER BY created_at ASC
        """,
        (day_start, day_end),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _fetch_expenses_for_shift(shift):
    if not shift or not shift.get("id"):
        return []

    end_at = shift.get("ended_at") or _now_str()
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense') AS name,
               COUNT(*) AS count,
               SUM(amount) AS total
        FROM manual_expenses
        WHERE shift_id = ?
           OR (
                (shift_id IS NULL OR TRIM(shift_id) = '')
                AND expense_at >= ?
                AND expense_at <= ?
           )
        GROUP BY COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense')
        ORDER BY name ASC
        """,
        (shift["id"], shift.get("started_at"), end_at),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _fetch_expenses_for_day(day_start, day_end):
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense') AS name,
               COUNT(*) AS count,
               SUM(amount) AS total
        FROM manual_expenses
        WHERE expense_at >= ? AND expense_at <= ?
        GROUP BY COALESCE(NULLIF(TRIM(note), ''), 'Manual Expense')
        ORDER BY name ASC
        """,
        (day_start, day_end),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _build_sales_summary(transactions, label, selected_date=None, active_shift=None):
    day_start = None
    day_end = None
    if selected_date:
        day_start, day_end = _operational_day_window(selected_date)

    if active_shift and active_shift.get("id"):
        job_orders = _fetch_job_orders_for_shift(active_shift)
        expense_rows = _fetch_expenses_for_shift(active_shift)
    elif selected_date:
        job_orders = _fetch_job_orders_for_day(day_start, day_end)
        expense_rows = _fetch_expenses_for_day(day_start, day_end)
    else:
        job_orders = []
        expense_rows = []

    valid_sales = [
        t for t in transactions
        if t.get("status") in ("COMPLETED", "SIMULATED")
    ]
    txn_ids = [t["id"] for t in valid_sales]
    job_order_ids = [t.get("job_order_id") for t in valid_sales if t.get("job_order_id")]
    
    item_usage = _fetch_item_usage(txn_ids)
    item_breakdown = _fetch_item_breakdown(txn_ids)

    gross_collected = sum(int(t.get("amount") or 0) for t in valid_sales)
    product_revenue = sum(int(t.get("product_total") or 0) for t in valid_sales)
    # Track transaction-native service revenue — these are embedded in gross_collected.
    # JO-sourced service items are NOT in gross_collected, so must not be used for
    # cash/net adjustment to avoid double-counting.
    txn_service_revenue = sum(int(t.get("service_total") or 0) for t in valid_sales)
    service_revenue = txn_service_revenue

    txn_gcash_revenue = 0
    txn_gcash_count = 0
    for t in valid_sales:
        amount = int(t.get("amount") or 0)
        if t.get("gcash_amount") is not None:
            gcash = int(t.get("gcash_amount") or 0)
        else:
            gcash = amount if int(t.get("paid_by_gcash") or 0) == 1 else 0
        gcash = max(0, min(gcash, amount))
        txn_gcash_revenue += gcash
        if gcash > 0:
            txn_gcash_count += 1

    post_cycle = {"count": 0, "amount": 0}
    if active_shift and active_shift.get("id"):
        post_cycle = summarize_post_cycle_payment_logs(
            shift_id=active_shift["id"],
            start_at=active_shift.get("started_at"),
            end_at=active_shift.get("ended_at") or _now_str(),
        )
    elif selected_date:
        post_cycle = summarize_post_cycle_payment_logs(start_at=day_start, end_at=day_end)

    post_cycle_amount = int(post_cycle.get("amount") or 0)
    post_cycle_count = int(post_cycle.get("count") or 0)

    total_sales = gross_collected
    machine_revenue = max(0, total_sales - product_revenue - service_revenue)

    gcash_revenue = txn_gcash_revenue + post_cycle_amount
    gcash_transaction_count = txn_gcash_count + post_cycle_count

    job_order_count = len(job_orders)
    job_order_used_count = sum(1 for r in job_orders if str(r.get("status") or "").upper() == "USED")
    job_order_open_count = sum(1 for r in job_orders if str(r.get("status") or "").upper() == "OPEN")
    job_order_total_amount = sum(int(r.get("total_amount") or 0) for r in job_orders)

    promo_orders = [r for r in job_orders if r.get("promo_id") or r.get("promo_name")]
    job_order_promo_count = len(promo_orders)
    promo_counts = {}
    for r in promo_orders:
        name = str(r.get("promo_name") or r.get("promo_id") or "Unknown Promo").strip()
        promo_counts[name] = promo_counts.get(name, 0) + 1
    job_order_promo_breakdown = [{"name": name, "count": count} for name, count in promo_counts.items()]

    manual_expenses = sum(int(r.get("total") or 0) for r in expense_rows)
    cogs_total = int(item_usage.get("cogs_total") or 0)
    total_expenses = manual_expenses + cogs_total

    analytics_settings = get_analytics_settings()
    include_service_in_net = bool(analytics_settings.get("include_service_revenue_in_net", False))
    _svc_adj = 0 if include_service_in_net else -txn_service_revenue

    cash_sales, cash_collected = _compute_cash_sales_and_net(
        total_sales,
        txn_gcash_revenue,
        post_cycle_amount,
        total_expenses,
        _svc_adj,
    )

    expense_breakdown = [
        {
            "name": (str(r.get("name") or "").strip() or "Manual Expense"),
            "count": int(r.get("count") or 0),
            "total": int(r.get("total") or 0),
        }
        for r in expense_rows
    ]
    if cogs_total > 0:
        expense_breakdown.append(
            {
                "name": "COGS",
                "count": int(item_usage.get("product_qty") or 0),
                "total": cogs_total,
            }
        )

    return {
        "label": label,
        "selected_date": selected_date,
        "transaction_count": len(valid_sales),
        "gross_collected": gross_collected,
        "total_sales": total_sales,
        "manual_expenses": manual_expenses,
        "cogs_total": cogs_total,
        "total_expenses": total_expenses,
        "net_sales": total_sales + _svc_adj - total_expenses,
        "net_after_expenses": total_sales + _svc_adj - total_expenses,
        "machine_revenue": machine_revenue,
        "cash_sales": cash_sales,
        "cash_revenue": cash_collected,
        "cash_collected": cash_collected,
        "gcash_revenue": gcash_revenue,
        "gcash_collected": gcash_revenue,
        "gcash_job_order_count": gcash_transaction_count,
        "post_cycle_transfer_amount": post_cycle_amount,
        "post_cycle_transfer_count": post_cycle_count,
        "product_revenue": product_revenue,
        "service_revenue": service_revenue,
        "product_qty": item_usage["product_qty"],
        "service_qty": item_usage["service_qty"],
        "product_breakdown": item_breakdown["product"],
        "service_breakdown": item_breakdown["service"],
        "expense_breakdown": expense_breakdown,
        "job_order_count": job_order_count,
        "job_order_used_count": job_order_used_count,
        "job_order_open_count": job_order_open_count,
        "job_order_total_amount": job_order_total_amount,
        "job_order_promo_count": job_order_promo_count,
        "job_order_promo_breakdown": job_order_promo_breakdown,
        "active_shift": active_shift,
    }


def _build_calendar_summary(year, month):
    year = int(year)
    month = int(month)
    days_in_month = calendar.monthrange(year, month)[1]
    month_key = f"{year:04d}-{month:02d}"

    month_start = f"{year:04d}-{month:02d}-01 00:00:00"
    month_end   = f"{year:04d}-{month:02d}-{days_in_month:02d} 23:59:59"

    conn = get_connection()

    # Pull transactions directly by their own started_at timestamp.
    all_tx_rows = conn.execute(
        """
        SELECT id, substr(started_at, 1, 10) AS tx_day, amount, service_total, status
        FROM transactions
        WHERE started_at >= ? AND started_at <= ?
          AND status IN ('COMPLETED', 'SIMULATED')
        """,
        (month_start, month_end),
    ).fetchall()

    # Pull COGS for the same set of transactions.
    all_tx_ids = [r["id"] for r in all_tx_rows]
    cogs_by_txn = {}
    if all_tx_ids:
        cogs_placeholders = ",".join("?" for _ in all_tx_ids)
        cogs_item_rows = conn.execute(
            f"""
            SELECT transaction_id, SUM(line_cost) AS cogs_total
            FROM transaction_items
            WHERE transaction_id IN ({cogs_placeholders})
            GROUP BY transaction_id
            """,
            all_tx_ids,
        ).fetchall()
        cogs_by_txn = {r["transaction_id"]: int(r["cogs_total"] or 0) for r in cogs_item_rows}

    expense_rows = conn.execute(
        """
        SELECT substr(expense_at, 1, 10) AS day,
               SUM(amount) AS manual_expenses
        FROM manual_expenses
        WHERE substr(expense_at, 1, 7) = ?
        GROUP BY substr(expense_at, 1, 10)
        ORDER BY day
        """,
        (month_key,),
    ).fetchall()
    conn.close()

    # Aggregate by the transaction's own recorded date.
    tx_map = {}
    for r in all_tx_rows:
        day = r["tx_day"]
        if not day or not day.startswith(month_key):
            continue

        entry = tx_map.setdefault(
            day,
            {
                "transaction_count": 0,
                "gross_sales": 0,
                "total_sales": 0,
                "cogs_total": 0,
            },
        )
        entry["transaction_count"] += 1
        amount = int(r["amount"] or 0)
        service_total = int(r["service_total"] or 0)
        entry["gross_sales"] += amount
        entry["total_sales"] += max(0, amount - service_total)
        entry["cogs_total"] += cogs_by_txn.get(r["id"], 0)

    expense_map = {row["day"]: int(row["manual_expenses"] or 0) for row in expense_rows}

    data = []
    for day in range(1, days_in_month + 1):
        key = f"{year:04d}-{month:02d}-{day:02d}"
        tx = tx_map.get(key, {})
        gross_sales = int(tx.get("gross_sales") or 0)
        total_sales = int(tx.get("total_sales") or 0)
        manual_expenses = int(expense_map.get(key) or 0)
        cogs_total = int(tx.get("cogs_total") or 0)
        total_expenses = manual_expenses + cogs_total

        entry = {
            "day": key,
            "transaction_count": int(tx.get("transaction_count") or 0),
            "gross_sales": gross_sales,
            "manual_expenses": manual_expenses,
            "cogs_total": cogs_total,
            "total_expenses": total_expenses,
            "net_sales": total_sales - total_expenses,
            "total_sales": total_sales,
        }
        data.append(entry)
    return data


@dashboard_bp.route("/")
def index():
    machines = _enrich_machines([_normalize_machine_runtime(m) for m in get_all_machines()])
    all_transactions = _enrich_transactions(_get_transactions(limit=None))
    transactions = all_transactions[:2000]

    stats = _build_dashboard_stats(all_transactions, machines)
    analytics = _build_analytics(all_transactions, machines)

    locations = [{
        "id": DEFAULT_LOCATION_ID,
        "name": "Local Pi",
        "pi_url": f"http://127.0.0.1:{os.environ.get('PORT', '5000')}",
    }]

    import json
    for m in machines:
        try:
            m["custom_modes_parsed"] = json.loads(m.get("custom_modes") or "[]")
        except Exception:
            m["custom_modes_parsed"] = []

    layout_json = get_dashboard_layout() or "null"

    return render_template(
        "dashboard.html",
        stats=stats,
        transactions=transactions,
        locations=locations,
        machines=machines,
        analytics=analytics,
        transactions_json=json.dumps(transactions),
        analytics_json=json.dumps(analytics),
        machines_json=json.dumps(machines),
        locations_json=json.dumps(locations),
        layout_json=layout_json,
    )


@dashboard_bp.route("/dashboard/analytics")
def get_analytics():
    start = request.args.get("start")
    end = request.args.get("end")

    start_date = datetime.strptime(start, "%Y-%m-%d").date() if start else None
    end_date = datetime.strptime(end, "%Y-%m-%d").date() if end else None

    machines = _enrich_machines(get_all_machines())
    transactions = _enrich_transactions(_get_transactions_for_date_range(start_date, end_date))
    filtered = _filter_by_date(transactions, start_date, end_date)

    return jsonify(_build_analytics(filtered, machines, start_date=start_date, end_date=end_date))


@dashboard_bp.route("/dashboard/post-cycle-payment/log", methods=["POST"])
def dashboard_log_post_cycle_payment():
    data = request.get_json(silent=True) or {}
    raw_amount = data.get("gcash_amount")
    note = str(data.get("note") or "").strip() or None
    location_id = str(data.get("location_id") or DEFAULT_LOCATION_ID).strip() or DEFAULT_LOCATION_ID

    try:
        amount = int(raw_amount)
    except (TypeError, ValueError):
        return jsonify({"error": "gcash_amount must be a whole number"}), 400

    if amount == 0:
        return jsonify({"error": "gcash_amount cannot be 0"}), 400

    active_shift = get_active_shift(location_id)
    shift_id = active_shift.get("id") if active_shift else None
    employee_id = active_shift.get("employee_id") if active_shift else None

    entry = create_post_cycle_payment_log(
        amount=amount,
        logged_at=_now_str(),
        shift_id=shift_id,
        employee_id=employee_id,
        note=note,
    )

    return jsonify(
        {
            "status": "ok",
            "warning": "No matching transaction was linked. Logged as a post-cycle transfer.",
            "message": "Post-cycle payment logged. Cash revenue reduced and GCash revenue increased.",
            "post_cycle_log": entry,
        }
    )


@dashboard_bp.route("/dashboard/receipt/override/save", methods=["POST"])
def dashboard_save_receipt_override():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}
    reference_id = str(data.get("reference_id") or "").strip()
    kind = str(data.get("kind") or "DAY").strip().upper()
    overrides = data.get("overrides")

    if not reference_id:
        return jsonify({"error": "reference_id is required"}), 400
    if kind not in ("DAY", "SHIFT"):
        return jsonify({"error": "kind must be DAY or SHIFT"}), 400
    if not isinstance(overrides, dict) or not overrides:
        return jsonify({"error": "overrides must be a non-empty object"}), 400

    save_receipt_overrides(reference_id, kind, overrides)
    return jsonify({"status": "ok", "reference_id": reference_id, "kind": kind})


@dashboard_bp.route("/dashboard/post-cycle-payment/add-ons", methods=["POST"])
def dashboard_log_post_cycle_addons():
    data = request.get_json(silent=True) or {}
    location_id = str(data.get("location_id") or DEFAULT_LOCATION_ID).strip() or DEFAULT_LOCATION_ID
    paid_by_gcash = bool(data.get("paid_by_gcash"))
    sale_items_raw = data.get("sale_items") if isinstance(data.get("sale_items"), list) else []

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400

    normalized_sale_items = []
    for item in sale_items_raw:
        kind = str((item or {}).get("kind") or "").strip().lower()
        item_id = str((item or {}).get("item_id") or "").strip()
        try:
            quantity = int((item or {}).get("quantity") or 0)
        except (TypeError, ValueError):
            quantity = 0

        if kind not in ("product", "service") or not item_id or quantity <= 0:
            continue

        normalized_sale_items.append(
            {
                "kind": kind,
                "item_id": item_id,
                "quantity": quantity,
            }
        )

    if not normalized_sale_items:
        return jsonify({"error": "Please add at least one product or service."}), 400

    timestamp = _now_str()
    txn_id = str(uuid.uuid4())
    request_id = "postcycle-addon-" + datetime.now().strftime("%Y%m%d%H%M%S") + "-" + txn_id[:8]

    txn_result = insert_transaction_with_items(
        txn_id=txn_id,
        machine_id="post-cycle-addons",
        base_amount=0,
        status="COMPLETED",
        started_at=timestamp,
        employee_id=active_shift.get("employee_id"),
        shift_id=active_shift.get("id"),
        sale_items=normalized_sale_items,
        request_id=request_id,
        paid_by_gcash=paid_by_gcash,
    )

    try_immediate_sync()

    total_amount = int(txn_result.get("total_amount") or 0)
    gcash_amount = total_amount if paid_by_gcash else 0
    cash_amount = total_amount - gcash_amount

    return jsonify(
        {
            "status": "ok",
            "message": "Added services/products logged successfully.",
            "transaction_id": str(txn_result.get("transaction_id") or txn_id),
            "total_amount": total_amount,
            "product_total": int(txn_result.get("product_total") or 0),
            "service_total": int(txn_result.get("service_total") or 0),
            "item_count": int(txn_result.get("item_count") or 0),
            "paid_by_gcash": 1 if paid_by_gcash else 0,
            "gcash_amount": gcash_amount,
            "cash_amount": cash_amount,
            "low_stock_warnings": txn_result.get("low_stock_warnings") or [],
        }
    ), 201


@dashboard_bp.route("/dashboard/summary/shift")
def get_current_shift_summary():
    location_id = request.args.get("location_id") or DEFAULT_LOCATION_ID
    active_shift = get_active_shift(location_id)

    if not active_shift:
        return jsonify(_build_sales_summary([], "Current Shift", active_shift=None))

    txns = _fetch_transactions_by_shift(active_shift["id"])
    return jsonify(_build_sales_summary(txns, "Current Shift", active_shift=active_shift))


@dashboard_bp.route("/dashboard/summary/shift/<shift_id>")
def get_specific_shift_summary(shift_id):
    shift = get_shift(shift_id)
    if not shift:
        return jsonify({"error": "Shift not found"}), 404

    txns = _fetch_transactions_by_shift(shift_id)
    return jsonify(_build_sales_summary(txns, "Shift Summary", active_shift=shift))


@dashboard_bp.route("/dashboard/summary/day")
def get_day_summary():
    date_arg = request.args.get("date")
    if date_arg:
        target_date = datetime.strptime(date_arg, "%Y-%m-%d").date()
    else:
        target_date = datetime.strptime(_current_operational_day_str(), "%Y-%m-%d").date()

    day_str = target_date.strftime("%Y-%m-%d")
    txns = _fetch_transactions_by_day(day_str)
    return jsonify(_build_sales_summary(txns, "Daily Summary", selected_date=day_str))


@dashboard_bp.route("/dashboard/settings/recent-shift-summary-count", methods=["GET"])
def get_recent_shift_summary_count_setting():
    return jsonify({"recent_shift_summary_count": get_recent_shift_summary_count()})


@dashboard_bp.route("/dashboard/admin/verify-pin", methods=["POST"])
def verify_admin_pin_setting():
    guard = require_admin_pin()
    if guard:
        return guard
    return jsonify({"status": "ok"}), 200


@dashboard_bp.route("/dashboard/settings/recent-shift-summary-count", methods=["POST"])
def update_recent_shift_summary_count_setting():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}
    try:
        updated = set_recent_shift_summary_count(data.get("recent_shift_summary_count"))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    return jsonify({"status": "ok", "recent_shift_summary_count": updated}), 200


@dashboard_bp.route("/dashboard/settings/receipt-format", methods=["GET"])
def get_receipt_format_setting():
    guard = require_admin_pin()
    if guard:
        return guard
    return jsonify(get_receipt_format_config())


@dashboard_bp.route("/dashboard/settings/receipt-format", methods=["POST"])
def update_receipt_format_setting():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}
    updated = set_receipt_format_config(data)
    return jsonify({"status": "ok", **updated}), 200


@dashboard_bp.route("/dashboard/settings/day-change-time", methods=["GET"])
def get_day_change_time_setting():
    return jsonify({"day_change_time": get_day_change_time()}), 200


@dashboard_bp.route("/dashboard/settings/day-change-time", methods=["POST"])
def update_day_change_time_setting():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}
    new_time = data.get("day_change_time")
    if not new_time:
        return jsonify({"error": "Missing day_change_time"}), 400

    try:
        set_day_change_time(new_time)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    return jsonify({"status": "ok", "day_change_time": get_day_change_time()}), 200


@dashboard_bp.route("/dashboard/settings/analytics", methods=["GET"])
def get_analytics_settings_route():
    guard = require_admin_pin()
    if guard:
        return guard
    return jsonify(get_analytics_settings()), 200


@dashboard_bp.route("/dashboard/settings/analytics", methods=["POST"])
def update_analytics_settings_route():
    guard = require_admin_pin()
    if guard:
        return guard
    data = request.get_json(silent=True) or {}
    try:
        updated = set_analytics_settings(data)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"status": "ok", **updated}), 200


@dashboard_bp.route("/dashboard/settings/layout", methods=["GET"])
def get_dashboard_layout_route():
    layout_data = get_dashboard_layout()
    import json
    try:
        parsed = json.loads(layout_data) if layout_data else None
    except Exception:
        parsed = None
    return jsonify({"layout": parsed}), 200


@dashboard_bp.route("/dashboard/settings/layout", methods=["POST"])
def update_dashboard_layout_route():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}
    layout_val = data.get("layout")
    if layout_val is None:
        return jsonify({"error": "Missing layout configuration"}), 400

    import json
    try:
        layout_str = json.dumps(layout_val)
        set_dashboard_layout(layout_str)
    except Exception as exc:
        return jsonify({"error": str(exc)}), 400

    return jsonify({"status": "ok", "layout": layout_val}), 200


@dashboard_bp.route("/dashboard/summary/calendar")
def get_calendar_summary():
    now = datetime.now()
    year = request.args.get("year", now.year)
    month = request.args.get("month", now.month)
    return jsonify({
        "year": int(year),
        "month": int(month),
        "days": _build_calendar_summary(year, month),
    })


@dashboard_bp.route("/dashboard/customers", methods=["GET"])
def dashboard_list_customers():
    limit_arg = request.args.get("limit", 500)
    try:
        limit = int(limit_arg)
    except (TypeError, ValueError):
        limit = 500
    return jsonify({"customers": list_customers(limit=limit)})


@dashboard_bp.route("/dashboard/customers", methods=["POST"])
def dashboard_create_or_update_customer():
    data = request.get_json(silent=True) or {}
    try:
        customer = _resolve_customer_payload({"customer": data})
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"status": "ok", "customer": customer}), 201


@dashboard_bp.route("/dashboard/job-orders/open", methods=["GET"])
def dashboard_list_open_job_orders():
    customer_id = request.args.get("customer_id")
    machine_id = request.args.get("machine_id")
    machine_type = request.args.get("machine_type")
    limit_arg = request.args.get("limit", 100)
    try:
        limit = int(limit_arg)
    except (TypeError, ValueError):
        limit = 100

    orders = list_open_job_orders(
        customer_id=customer_id,
        machine_id=machine_id,
        machine_type=machine_type,
        limit=limit,
    )
    return jsonify({"job_orders": orders})


@dashboard_bp.route("/dashboard/job-orders", methods=["POST"])
def dashboard_create_job_order():
    data = request.get_json(silent=True) or {}
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID

    machine_orders_payload = data.get("machine_orders") if isinstance(data.get("machine_orders"), list) else None

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400
    try:
        customer = _resolve_customer_payload(data)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    wash_qty = 0
    dry_qty = 0
    product_qty = 0
    service_qty = 0
    product_amount = 0
    service_amount = 0
    paid_by_gcash = False
    wash_mode = "normal"
    dry_mode = "normal"
    promo_id = None
    promo_name = None
    print_receipt = 1
    items = []

    if not machine_orders_payload:
        try:
            wash_qty = int(data.get("wash_qty", 0))
            dry_qty = int(data.get("dry_qty", 0))
            product_qty = int(data.get("product_qty", 0))
            service_qty = int(data.get("service_qty", 0))
            product_amount = int(data.get("product_amount", 0))
            service_amount = int(data.get("service_amount", 0))
        except (TypeError, ValueError):
            return jsonify({"error": "wash_qty, dry_qty, product_qty, service_qty, product_amount, and service_amount must be whole numbers"}), 400

        paid_by_gcash = bool(data.get("paid_by_gcash"))
        wash_mode = _normalize_job_order_mode(data.get("wash_mode"))
        dry_mode = _normalize_job_order_mode(data.get("dry_mode"))
        promo_id = str(data.get("promo_id") or "").strip() or None
        promo_name = str(data.get("promo_name") or "").strip() or None
        print_receipt = 1 if data.get("print_receipt", True) else 0
        items = data.get("items") or []

        if wash_qty < 0 or dry_qty < 0 or product_qty < 0 or service_qty < 0 or product_amount < 0 or service_amount < 0:
            return jsonify({"error": "wash_qty, dry_qty, product_qty, service_qty, product_amount, and service_amount cannot be negative"}), 400
        if wash_qty == 0 and dry_qty == 0:
            return jsonify({"error": "At least one of wash_qty or dry_qty must be greater than zero"}), 400

    try:
        if machine_orders_payload:
            normalized_orders = []
            for entry in machine_orders_payload:
                machine_id = str((entry or {}).get("machine_id") or "").strip()
                machine_type = str((entry or {}).get("machine_type") or "").strip().lower()

                machine = None
                if machine_id:
                    machine = get_machine(machine_id)
                    if not machine:
                        return jsonify({"error": f"Machine not found: {machine_id}"}), 404
                    machine_type = str(machine.get("type") or "").strip().lower()
                elif machine_type not in ("washer", "dryer"):
                    return jsonify({"error": "Each machine order requires machine_id or valid machine_type (washer/dryer)"}), 400

                try:
                    entry_wash_qty = int((entry or {}).get("wash_qty", 0))
                    entry_dry_qty = int((entry or {}).get("dry_qty", 0))
                    entry_product_qty = int((entry or {}).get("product_qty", 0))
                    entry_service_qty = int((entry or {}).get("service_qty", 0))
                    entry_product_amount = int((entry or {}).get("product_amount", 0))
                    entry_service_amount = int((entry or {}).get("service_amount", 0))
                except (TypeError, ValueError):
                    return jsonify({"error": "machine_orders wash_qty, dry_qty, product_qty, service_qty, product_amount, and service_amount must be whole numbers"}), 400

                paid_by_gcash = paid_by_gcash or bool((entry or {}).get("paid_by_gcash"))
                entry_wash_mode = _normalize_job_order_mode((entry or {}).get("wash_mode"))
                entry_dry_mode = _normalize_job_order_mode((entry or {}).get("dry_mode"))
                entry_promo_id = str((entry or {}).get("promo_id") or "").strip() or None
                entry_promo_name = str((entry or {}).get("promo_name") or "").strip() or None
                entry_print = 1 if (entry or {}).get("print_receipt", True) else 0
                entry_items = (entry or {}).get("items") or []

                if entry_wash_qty < 0 or entry_dry_qty < 0 or entry_product_qty < 0 or entry_service_qty < 0 or entry_product_amount < 0 or entry_service_amount < 0:
                    return jsonify({"error": "machine_orders wash_qty, dry_qty, product_qty, service_qty, product_amount, and service_amount cannot be negative"}), 400
                if entry_wash_qty == 0 and entry_dry_qty == 0:
                    return jsonify({"error": "machine_orders entries must include wash_qty or dry_qty"}), 400

                wash_unit_price = _price_by_machine_type("washer", entry_wash_mode, machine) if entry_wash_qty > 0 else _price_by_machine_type("washer", "standard")
                dry_unit_price = _price_by_machine_type("dryer", entry_dry_mode, machine) if entry_dry_qty > 0 else _price_by_machine_type("dryer", "standard")

                normalized_orders.append({
                    "machine_id": machine_id or ("any-washer" if machine_type == "washer" else "any-dryer"),
                    "machine_name": str((machine or {}).get("name") or ("Any Washer" if machine_type == "washer" else "Any Dryer")),
                    "machine_type": machine_type,
                    "wash_mode": entry_wash_mode,
                    "dry_mode": entry_dry_mode,
                    "wash_qty": entry_wash_qty,
                    "dry_qty": entry_dry_qty,
                    "product_qty": entry_product_qty,
                    "service_qty": entry_service_qty,
                    "product_amount": entry_product_amount,
                    "service_amount": entry_service_amount,
                    "paid_by_gcash": paid_by_gcash or bool((entry or {}).get("paid_by_gcash")),
                    "wash_unit_price": wash_unit_price,
                    "dry_unit_price": dry_unit_price,
                    "promo_id": entry_promo_id,
                    "promo_name": entry_promo_name,
                    "print_receipt": entry_print,
                    "items": entry_items,
                })

            created_orders = create_job_orders_bulk(
                customer_id=customer["customer_id"],
                customer_name=customer["name"],
                customer_phone=customer.get("phone"),
                machine_orders=normalized_orders,
                created_by_shift_id=active_shift.get("id"),
                created_by_employee_id=active_shift.get("employee_id"),
                created_by_employee_name=active_shift.get("display_name"),
            )
            return jsonify({"status": "ok", "job_orders": created_orders, "created_count": len(created_orders)}), 201

        machine_id = str(data.get("machine_id") or "").strip()
        machine = get_machine(machine_id) if machine_id else None

        machine_type = "mixed" if wash_qty > 0 and dry_qty > 0 else ("washer" if wash_qty > 0 else "dryer")
        machine_name = "Grouped Washer + Dryer" if machine_type == "mixed" else (
            "Any Washer" if machine_type == "washer" else "Any Dryer"
        )
        if machine:
            machine_type = str(machine.get("type") or "").strip().lower()
            machine_name = str(machine.get("name") or machine_id)
        elif machine_id:
            return jsonify({"error": "Machine not found"}), 404

        wash_unit_price = _price_by_machine_type("washer", wash_mode, machine) if wash_qty > 0 else _price_by_machine_type("washer", "standard")
        dry_unit_price = _price_by_machine_type("dryer", dry_mode, machine) if dry_qty > 0 else _price_by_machine_type("dryer", "standard")

        order = create_job_order(
            customer_id=customer["customer_id"],
            customer_name=customer["name"],
            customer_phone=customer.get("phone"),
            machine_id=(machine_id if machine_id else ("any-mixed" if machine_type == "mixed" else ("any-washer" if machine_type == "washer" else "any-dryer"))),
            machine_name=machine_name,
            machine_type=machine_type,
            wash_mode=wash_mode,
            dry_mode=dry_mode,
            wash_qty=wash_qty,
            dry_qty=dry_qty,
            wash_unit_price=wash_unit_price,
            dry_unit_price=dry_unit_price,
            product_qty=product_qty,
            service_qty=service_qty,
            product_amount=product_amount,
            service_amount=service_amount,
            paid_by_gcash=paid_by_gcash,
            created_by_shift_id=active_shift.get("id"),
            created_by_employee_id=active_shift.get("employee_id"),
            created_by_employee_name=active_shift.get("display_name"),
            promo_id=promo_id,
            promo_name=promo_name,
            print_receipt=print_receipt,
            items=items,
        )
        return jsonify({"status": "ok", "job_order": order}), 201

    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": "Failed to create job order"}), 500


@dashboard_bp.route("/dashboard/job-orders/<job_order_id>", methods=["DELETE"])
def dashboard_delete_job_order(job_order_id):
    order = get_job_order(job_order_id)
    if not order:
        return jsonify({"error": "Job order not found"}), 404

    status = str(order.get("status") or "").strip().upper()
    if status != "OPEN":
        return jsonify({"error": "Only open job orders can be deleted"}), 409

    if not delete_job_order(job_order_id):
        return jsonify({"error": "Failed to delete job order"}), 500

    return jsonify({"status": "ok", "deleted_job_order_id": job_order_id})


@dashboard_bp.route("/dashboard/machine/start", methods=["POST"])
@_json_api_guard
def dashboard_start_machine():
    data = request.get_json() or {}
    machine_id = data.get("machine_id")
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID
    sale_items = _exclude_quick_service_items(data.get("sale_items") or [])
    request_id = data.get("request_id")
    job_order_id = data.get("job_order_id")
    activation_mode = str(data.get("activation_mode") or "standard").strip().lower()
    paid_by_gcash = bool(data.get("paid_by_gcash"))

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    if request_id:
        existing = get_transaction_by_request_id(request_id)
        if existing:
            existing_machine = get_machine(machine_id)
            runtime_machine = _normalize_machine_runtime(existing_machine) if existing_machine else {}
            return jsonify({
                "status": existing["status"],
                "transaction_id": existing["id"],
                "stored_transaction_id": existing["id"],
                "machine": machine["name"],
                "amount": int(existing.get("amount") or 0),
                "base_amount": int(existing.get("amount") or 0) - int(existing.get("product_total") or 0) - int(existing.get("service_total") or 0),
                "product_total": int(existing.get("product_total") or 0),
                "service_total": int(existing.get("service_total") or 0),
                "item_count": int(existing.get("item_count") or 0),
                "low_stock_warnings": [],
                "idempotent_hit": True,
                "remaining_seconds": int(runtime_machine.get("remaining_seconds") or 0),
                "employee_id": existing.get("employee_id"),
                "shift_id": existing.get("shift_id"),
                "customer_id": existing.get("customer_id"),
                "customer_name": existing.get("customer_name"),
                "customer_phone": existing.get("customer_phone"),
                "paid_by_gcash": int(existing.get("paid_by_gcash") or 0),
                "job_order_id": existing.get("job_order_id"),
                "job_order_no": existing.get("job_order_no"),
            })

    if machine.get("status") == "BUSY":
        return jsonify({"error": "Machine is already running"}), 409

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400

    if not str(job_order_id or "").strip():
        customer_payload = None
        customer_data = data.get("customer") if isinstance(data.get("customer"), dict) else None
        if customer_data:
            customer_name = str(customer_data.get("name") or "").strip()
            if customer_name:
                try:
                    customer_payload = _resolve_customer_payload({"customer": customer_data})
                except ValueError:
                    customer_payload = None

        try:
            activation_result = _activate_machine_with_sale(
                machine=machine,
                active_shift=active_shift,
                sale_items=sale_items,
                request_id=request_id,
                customer_payload=customer_payload,
                paid_by_gcash=paid_by_gcash,
                activation_mode=activation_mode,
            )
        except RuntimeError as exc:
            return jsonify({"error": str(exc), "status": "ERROR"}), 502

        try_immediate_sync()
        return jsonify(activation_result)

    requested_order = get_job_order(job_order_id)
    if not requested_order:
        return jsonify({"error": "Job order not found"}), 404
    if str(requested_order.get("status") or "").upper() != "OPEN":
        return jsonify({"error": "Job order is already used or closed"}), 409
    requested_machine_type = str(requested_order.get("machine_type") or "").strip().lower()
    current_machine_type = str(machine.get("type") or "").strip().lower()
    if requested_machine_type not in ("mixed", current_machine_type):
        return jsonify({"error": "Job order does not match selected machine type"}), 400

    machine_type = current_machine_type
    wash_mode = _normalize_job_order_mode(requested_order.get("wash_mode"))
    dry_mode = _normalize_job_order_mode(requested_order.get("dry_mode"))
    selected_mode = wash_mode if machine_type == "washer" else dry_mode

    if selected_mode == "quick":
        pulse_count = max(1, int(machine.get("quick_wash_pulse_count") or 1))
        base_amount = max(0, int(machine.get("quick_wash_price") or machine.get("vend_price") or 0))
    else:
        pulse_count = int(machine["pulse_count"])
        base_amount = max(0, int(machine.get("vend_price") or 0))

    timestamp = _now_str()

    import uuid

    txn_id = str(uuid.uuid4())
    txn_status = "COMPLETED"

    try:
        claimed_job_order = claim_job_order_for_activation(
            job_order_id=job_order_id,
            machine_id=machine_id,
            machine_type=machine.get("type"),
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), _job_order_error_status(str(exc))

    customer = {
        "customer_id": claimed_job_order.get("customer_id"),
        "name": claimed_job_order.get("customer_name"),
        "phone": claimed_job_order.get("customer_phone"),
    }

    txn_result = {
        "total_amount": 0,
        "product_total": 0,
        "service_total": 0,
        "item_count": 0,
        "low_stock_warnings": [],
        "idempotent_hit": False,
    }
    transaction_id = None
    is_job_order_completed = str(claimed_job_order.get("status") or "").upper() == "USED"
    target_amount = base_amount
    fetch_jo_items = False

    finalize_request_id = request_id or f"jo-act-{claimed_job_order.get('id')}:{txn_id}"
    txn_result = insert_transaction_with_items(
        txn_id,
        machine_id,
        target_amount,
        txn_status,
        timestamp,
        employee_id=active_shift["employee_id"],
        shift_id=active_shift["id"],
        sale_items=[],
        request_id=finalize_request_id,
        customer_id=customer["customer_id"],
        customer_name=customer["name"],
        customer_phone=customer.get("phone"),
        paid_by_gcash=bool(claimed_job_order.get("paid_by_gcash")),
        job_order_id=claimed_job_order.get("id"),
        job_order_no=claimed_job_order.get("job_order_no"),
        fetch_jo_items=fetch_jo_items,
    )
    transaction_id = txn_result.get("transaction_id", txn_id)

    if is_job_order_completed:
        conn = get_connection()
        row_sum = conn.execute(
            "SELECT SUM(amount) AS total FROM transactions WHERE job_order_id = ? AND status IN ('COMPLETED', 'SIMULATED')",
            (claimed_job_order.get("id"),)
        ).fetchone()
        already_paid = int(row_sum["total"] or 0) if row_sum else 0
        conn.close()

        addon_amount = int(claimed_job_order.get("total_amount") or 0) - already_paid
        has_items = bool(int(claimed_job_order.get("product_qty") or 0) > 0 or int(claimed_job_order.get("service_qty") or 0) > 0)

        if addon_amount != 0 or has_items:
            addon_txn_id = str(uuid.uuid4())
            addon_request_id = f"jo-addons-{claimed_job_order.get('id')}:{addon_txn_id[:8]}"
            insert_transaction_with_items(
                addon_txn_id,
                "post-cycle-addons",
                addon_amount,
                txn_status,
                timestamp,
                employee_id=active_shift["employee_id"],
                shift_id=active_shift["id"],
                sale_items=[],
                request_id=addon_request_id,
                customer_id=customer["customer_id"],
                customer_name=customer["name"],
                customer_phone=customer.get("phone"),
                paid_by_gcash=bool(claimed_job_order.get("paid_by_gcash")),
                job_order_id=claimed_job_order.get("id"),
                job_order_no=claimed_job_order.get("job_order_no"),
                fetch_jo_items=True,
            )
        attach_job_order_transaction(claimed_job_order.get("id"), transaction_id)

    updated_customer = increment_customer_order_count(customer["customer_id"], machine.get("type"), quantity=1) or customer
    run_ends_at = (datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S") + timedelta(seconds=MACHINE_RUN_SECONDS)).strftime("%Y-%m-%d %H:%M:%S")
    set_machine_run_window(machine_id, timestamp, run_ends_at)

    try_immediate_sync()

    def on_result(success, message):
        if not success and not IS_DEV:
            if transaction_id:
                delete_transactions([transaction_id])
            revert_job_order_usage(claimed_job_order.get("id"), machine.get("type"))
            clear_machine_run_window(machine_id)
            update_machine_status(machine_id, "OFFLINE")

    async_send_pulse(
        machine["esp32_ip"],
        machine["pulse_on"],
        machine["pulse_off"],
        pulse_count,
        on_result
    )


    return jsonify({
        "status": txn_status,
        "transaction_id": transaction_id,
        "stored_transaction_id": transaction_id,
        "machine": machine["name"],
        "amount": txn_result["total_amount"],
        "base_amount": base_amount,
        "activation_mode": selected_mode,
        "product_total": txn_result["product_total"],
        "service_total": txn_result["service_total"],
        "item_count": txn_result["item_count"],
        "low_stock_warnings": txn_result["low_stock_warnings"],
        "idempotent_hit": txn_result.get("idempotent_hit", False),
        "remaining_seconds": MACHINE_RUN_SECONDS,
        "employee_id": active_shift["employee_id"],
        "employee_name": active_shift.get("display_name"),
        "shift_id": active_shift["id"],
        "customer": updated_customer,
        "customer_id": updated_customer.get("customer_id"),
        "customer_name": updated_customer.get("name"),
        "customer_phone": updated_customer.get("phone"),
        "paid_by_gcash": int(txn_result.get("paid_by_gcash") or int(claimed_job_order.get("paid_by_gcash") or 0)),
        "job_order_id": claimed_job_order.get("id"),
        "job_order_no": claimed_job_order.get("job_order_no"),
        "job_order_status": claimed_job_order.get("status"),
        "job_order_remaining_wash_qty": int(claimed_job_order.get("wash_qty") or 0),
        "job_order_remaining_dry_qty": int(claimed_job_order.get("dry_qty") or 0),
    })


@dashboard_bp.route("/dashboard/machine/start-bulk", methods=["POST"])
@_json_api_guard
def dashboard_start_bulk_machines():
    data = request.get_json() or {}
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID
    sale_items = _exclude_quick_service_items(data.get("sale_items") or [])
    paid_by_gcash = bool(data.get("paid_by_gcash"))
    activation_mode = str(data.get("activation_mode") or "standard").strip().lower()
    request_group_id = str(data.get("request_id") or "").strip()

    machine_ids_raw = data.get("machine_ids") if isinstance(data.get("machine_ids"), list) else []
    machine_ids = []
    seen = set()
    for machine_id in machine_ids_raw:
        normalized = str(machine_id or "").strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        machine_ids.append(normalized)

    if not machine_ids:
        return jsonify({"error": "At least one machine must be selected"}), 400

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400

    customer_payload = None
    customer_data = data.get("customer") if isinstance(data.get("customer"), dict) else None
    if customer_data:
        customer_name = str(customer_data.get("name") or "").strip()
        if customer_name:
            try:
                customer_payload = _resolve_customer_payload({"customer": customer_data})
            except ValueError:
                customer_payload = None

    if not request_group_id:
        request_group_id = f"bulk-{int(datetime.now().timestamp())}"

    started = []
    failed = []
    apply_sale_items = list(sale_items)

    for idx, machine_id in enumerate(machine_ids):
        machine = get_machine(machine_id)
        if not machine:
            failed.append({"machine_id": machine_id, "error": "Machine not found"})
            continue

        if machine.get("status") == "BUSY":
            failed.append({"machine_id": machine_id, "machine": machine.get("name") or machine_id, "error": "Machine is already running"})
            continue

        if machine.get("status") == "OFFLINE":
            failed.append({"machine_id": machine_id, "machine": machine.get("name") or machine_id, "error": "Machine is offline"})
            continue

        scoped_request_id = f"{request_group_id}:{idx}:{machine_id}"

        try:
            result = _activate_machine_with_sale(
                machine=machine,
                active_shift=active_shift,
                sale_items=apply_sale_items,
                request_id=scoped_request_id,
                customer_payload=customer_payload,
                paid_by_gcash=paid_by_gcash,
                activation_mode=activation_mode,
            )
            started.append(result)
            apply_sale_items = []
        except RuntimeError as exc:
            failed.append({"machine_id": machine_id, "machine": machine.get("name") or machine_id, "error": str(exc)})

    if started:
        try_immediate_sync()

    if started and failed:
        status = "PARTIAL"
    elif started:
        status = "COMPLETED"
    else:
        status = "FAILED"

    response = {
        "status": status,
        "request_id": request_group_id,
        "started_count": len(started),
        "failed_count": len(failed),
        "started": started,
        "failed": failed,
        "transaction_ids": [row.get("stored_transaction_id") for row in started if row.get("stored_transaction_id")],
        "total_amount": sum(int(row.get("amount") or 0) for row in started),
        "remaining_seconds": MACHINE_RUN_SECONDS,
    }

    http_status = 200 if started else 409
    return jsonify(response), http_status


@dashboard_bp.route("/dashboard/machine/activate", methods=["POST"])
def dashboard_activate_machine_without_transaction():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json() or {}
    machine_id = data.get("machine_id")
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID

    try:
        pulse_count = int(data.get("pulse_count") or 1)
    except (TypeError, ValueError):
        return jsonify({"error": "pulse_count must be a valid number"}), 400

    if pulse_count < 1:
        return jsonify({"error": "pulse_count must be at least 1"}), 400

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    if machine.get("status") == "BUSY":
        return jsonify({"error": "Machine is already running"}), 409

    timestamp = _now_str()
    pulse_duration_ms = 50

    def on_result(success, message):
        if not success and not IS_DEV:
            clear_machine_run_window(machine_id)
            update_machine_status(machine_id, "OFFLINE")

    async_send_pulse(
        machine["esp32_ip"],
        pulse_duration_ms,
        pulse_duration_ms,
        pulse_count,
        on_result
    )

    run_ends_at = (datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S") + timedelta(seconds=MACHINE_RUN_SECONDS)).strftime("%Y-%m-%d %H:%M:%S")
    set_machine_run_window(machine_id, timestamp, run_ends_at)

    return jsonify(
        {
            "status": "COMPLETED",
            "recorded": False,
            "machine": machine["name"],
            "machine_id": machine_id,
            "location_id": location_id,
            "amount": 0,
            "base_amount": 0,
            "product_total": 0,
            "service_total": 0,
            "item_count": 0,
            "remaining_seconds": MACHINE_RUN_SECONDS,
            "pulse_count": pulse_count,
            "pulse_on": pulse_duration_ms,
            "pulse_off": pulse_duration_ms,
            "message": "Machine activated without transaction recording.",
        }
    )


@dashboard_bp.route("/dashboard/machine/quick-service", methods=["POST"])
def dashboard_quick_service():
    data = request.get_json() or {}
    machine_id = data.get("machine_id")
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID
    request_id = data.get("request_id")

    try:
        quantity = int(data.get("quantity", 1))
    except (TypeError, ValueError):
        return jsonify({"error": "quantity must be a valid number"}), 400

    if quantity < 1:
        return jsonify({"error": "quantity must be at least 1"}), 400

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404
    if machine.get("status") == "BUSY":
        return jsonify({"error": "Machine is already running"}), 409

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400

    service_name = _quick_service_name_for_machine(machine.get("type"))
    if not service_name:
        return jsonify({"error": "Quick service is only available for washer/dryer machines"}), 400

    service = get_active_service_by_name(service_name)
    if not service:
        return jsonify({"error": f"Service '{service_name}' is not configured or active"}), 404

    pulse_bonus = max(1, int(service.get("bonus_pulses") or 1))
    pulse_count = quantity * pulse_bonus
    timestamp = _now_str()

    import uuid

    txn_id = str(uuid.uuid4())
    txn_status = "COMPLETED"
    sale_items = [{"kind": "service", "item_id": service["id"], "quantity": quantity}]

    txn_result = insert_transaction_with_items(
        txn_id,
        machine_id,
        0,
        txn_status,
        timestamp,
        employee_id=active_shift["employee_id"],
        shift_id=active_shift["id"],
        sale_items=sale_items,
        request_id=request_id,
    )

    try_immediate_sync()

    final_txn_id = txn_result.get("transaction_id", txn_id)
    def on_result(success, message):
        if not success and not IS_DEV:
            delete_transactions([final_txn_id])
            clear_machine_run_window(machine_id)
            update_machine_status(machine_id, "OFFLINE")

    async_send_pulse(
        machine["esp32_ip"],
        machine["pulse_on"],
        machine["pulse_off"],
        pulse_count,
        on_result
    )

    return jsonify(
        {
            "status": txn_status,
            "mode": "quick_service",
            "machine": machine["name"],
            "machine_type": machine.get("type"),
            "service_name": service["name"],
            "quantity": quantity,
            "bonus_pulses": pulse_bonus,
            "pulse_count": pulse_count,
            "transaction_id": txn_result.get("transaction_id", txn_id),
            "service_total": txn_result["service_total"],
            "amount": txn_result["total_amount"],
            "employee_id": active_shift["employee_id"],
            "shift_id": active_shift["id"],
        }
    )


@dashboard_bp.route("/dashboard/machine/quick-wash", methods=["POST"])
@_json_api_guard
def dashboard_quick_wash():
    data = request.get_json() or {}
    machine_id = data.get("machine_id")
    location_id = data.get("location_id") or DEFAULT_LOCATION_ID
    request_id = data.get("request_id")
    job_order_id = data.get("job_order_id")

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    if str(machine.get("type") or "").strip().lower() != "washer":
        return jsonify({"error": "Quick wash is only available for washer machines"}), 400

    if request_id:
        existing = get_transaction_by_request_id(request_id)
        if existing:
            existing_machine = get_machine(machine_id)
            runtime_machine = _normalize_machine_runtime(existing_machine) if existing_machine else {}
            return jsonify(
                {
                    "status": existing["status"],
                    "mode": "quick_wash",
                    "transaction_id": existing["id"],
                    "stored_transaction_id": existing["id"],
                    "machine": machine["name"],
                    "machine_id": machine_id,
                    "amount": int(existing.get("amount") or 0),
                    "base_amount": int(existing.get("amount") or 0) - int(existing.get("product_total") or 0) - int(existing.get("service_total") or 0),
                    "quick_wash_price": int(existing.get("amount") or 0) - int(existing.get("product_total") or 0) - int(existing.get("service_total") or 0),
                    "quick_wash_pulse_count": max(1, int(machine.get("quick_wash_pulse_count") or 1)),
                    "product_total": int(existing.get("product_total") or 0),
                    "service_total": int(existing.get("service_total") or 0),
                    "item_count": int(existing.get("item_count") or 0),
                    "low_stock_warnings": [],
                    "idempotent_hit": True,
                    "remaining_seconds": int(runtime_machine.get("remaining_seconds") or 0),
                    "employee_id": existing.get("employee_id"),
                    "shift_id": existing.get("shift_id"),
                    "customer_id": existing.get("customer_id"),
                    "customer_name": existing.get("customer_name"),
                    "customer_phone": existing.get("customer_phone"),
                    "job_order_id": existing.get("job_order_id"),
                    "job_order_no": existing.get("job_order_no"),
                }
            )

    if machine.get("status") == "BUSY":
        return jsonify({"error": "Machine is already running"}), 409

    active_shift = get_active_shift(location_id)
    if not active_shift:
        return jsonify({"error": "No active employee shift. Please time in first."}), 400

    if not str(job_order_id or "").strip():
        return jsonify({"error": "Valid job_order_id is required before machine activation"}), 400

    timestamp = _now_str()
    pulse_count = max(1, int(machine.get("quick_wash_pulse_count") or 1))
    quick_wash_price = max(0, int(machine.get("quick_wash_price") or machine.get("vend_price") or 0))
    import uuid

    txn_id = str(uuid.uuid4())
    txn_status = "COMPLETED"

    try:
        claimed_job_order = claim_job_order_for_activation(
            job_order_id=job_order_id,
            machine_id=machine_id,
            machine_type=machine.get("type"),
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), _job_order_error_status(str(exc))

    customer = {
        "customer_id": claimed_job_order.get("customer_id"),
        "name": claimed_job_order.get("customer_name"),
        "phone": claimed_job_order.get("customer_phone"),
    }

    txn_result = {
        "total_amount": 0,
        "product_total": 0,
        "service_total": 0,
        "item_count": 0,
        "low_stock_warnings": [],
        "idempotent_hit": False,
    }
    transaction_id = None
    is_job_order_completed = str(claimed_job_order.get("status") or "").upper() == "USED"
    if is_job_order_completed:
        conn = get_connection()
    target_amount = quick_wash_price
    fetch_jo_items = False

    finalize_request_id = request_id or f"jo-act-{claimed_job_order.get('id')}:{txn_id}"
    txn_result = insert_transaction_with_items(
        txn_id,
        machine_id,
        target_amount,
        txn_status,
        timestamp,
        employee_id=active_shift["employee_id"],
        shift_id=active_shift["id"],
        sale_items=[],
        request_id=finalize_request_id,
        customer_id=customer["customer_id"],
        customer_name=customer["name"],
        customer_phone=customer.get("phone"),
        job_order_id=claimed_job_order.get("id"),
        job_order_no=claimed_job_order.get("job_order_no"),
        paid_by_gcash=bool(int(claimed_job_order.get("paid_by_gcash") or 0)),
        fetch_jo_items=fetch_jo_items,
    )
    transaction_id = txn_result.get("transaction_id", txn_id)

    if is_job_order_completed:
        conn = get_connection()
        row_sum = conn.execute(
            "SELECT SUM(amount) AS total FROM transactions WHERE job_order_id = ? AND status IN ('COMPLETED', 'SIMULATED')",
            (claimed_job_order.get("id"),)
        ).fetchone()
        already_paid = int(row_sum["total"] or 0) if row_sum else 0
        conn.close()

        addon_amount = int(claimed_job_order.get("total_amount") or 0) - already_paid
        has_items = bool(int(claimed_job_order.get("product_qty") or 0) > 0 or int(claimed_job_order.get("service_qty") or 0) > 0)

        if addon_amount != 0 or has_items:
            addon_txn_id = str(uuid.uuid4())
            addon_request_id = f"jo-addons-{claimed_job_order.get('id')}:{addon_txn_id[:8]}"
            insert_transaction_with_items(
                addon_txn_id,
                "post-cycle-addons",
                addon_amount,
                txn_status,
                timestamp,
                employee_id=active_shift["employee_id"],
                shift_id=active_shift["id"],
                sale_items=[],
                request_id=addon_request_id,
                customer_id=customer["customer_id"],
                customer_name=customer["name"],
                customer_phone=customer.get("phone"),
                job_order_id=claimed_job_order.get("id"),
                job_order_no=claimed_job_order.get("job_order_no"),
                paid_by_gcash=bool(int(claimed_job_order.get("paid_by_gcash") or 0)),
                fetch_jo_items=True,
            )
        attach_job_order_transaction(claimed_job_order.get("id"), transaction_id)

    updated_customer = increment_customer_order_count(customer["customer_id"], machine.get("type"), quantity=1) or customer

    run_ends_at = (datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S") + timedelta(seconds=MACHINE_RUN_SECONDS)).strftime(
        "%Y-%m-%d %H:%M:%S"
    )
    set_machine_run_window(machine_id, timestamp, run_ends_at)

    try_immediate_sync()

    def on_result(success, message):
        if not success and not IS_DEV:
            if transaction_id:
                delete_transactions([transaction_id])
            revert_job_order_usage(claimed_job_order.get("id"), machine.get("type"))
            clear_machine_run_window(machine_id)
            update_machine_status(machine_id, "OFFLINE")

    async_send_pulse(
        machine["esp32_ip"],
        machine["pulse_on"],
        machine["pulse_off"],
        pulse_count,
        on_result
    )

    return jsonify(
        {
            "status": txn_status,
            "mode": "quick_wash",
            "transaction_id": transaction_id,
            "stored_transaction_id": transaction_id,
            "machine": machine["name"],
            "machine_id": machine_id,
            "amount": txn_result["total_amount"],
            "base_amount": quick_wash_price,
            "quick_wash_price": quick_wash_price,
            "quick_wash_pulse_count": pulse_count,
            "product_total": txn_result["product_total"],
            "service_total": txn_result["service_total"],
            "item_count": txn_result["item_count"],
            "low_stock_warnings": txn_result["low_stock_warnings"],
            "idempotent_hit": txn_result.get("idempotent_hit", False),
            "remaining_seconds": MACHINE_RUN_SECONDS,
            "employee_id": active_shift["employee_id"],
            "employee_name": active_shift.get("display_name"),
            "shift_id": active_shift["id"],
            "customer": updated_customer,
            "customer_id": updated_customer.get("customer_id"),
            "customer_name": updated_customer.get("name"),
            "customer_phone": updated_customer.get("phone"),
            "job_order_id": claimed_job_order.get("id"),
            "job_order_no": claimed_job_order.get("job_order_no"),
            "job_order_status": claimed_job_order.get("status"),
            "job_order_remaining_wash_qty": int(claimed_job_order.get("wash_qty") or 0),
            "job_order_remaining_dry_qty": int(claimed_job_order.get("dry_qty") or 0),
        }
    )


@dashboard_bp.route("/dashboard/machine/stop", methods=["POST"])
def dashboard_stop_machine():
    data = request.get_json() or {}
    machine_id = data.get("machine_id")

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    clear_machine_run_window(machine_id)
    print(f"[{_now_str()}] Machine {machine['name']} stopped manually from dashboard")

    return jsonify({"status": "STOPPED", "machine": machine["name"]})


@dashboard_bp.route("/dashboard/machine/life", methods=["POST"])
def dashboard_machine_life_check():
    data = request.get_json() or {}
    machine_id = data.get("machine_id")

    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    ok, message = check_esp32_life(machine["esp32_ip"])
    if ok:
        return jsonify({"status": "ALIVE", "machine": machine["name"], "message": message}), 200
    return jsonify({"status": "OFFLINE", "machine": machine["name"], "error": message}), 502


@dashboard_bp.route("/dashboard/machine/settings", methods=["POST"])
def update_machine_settings():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json() or {}
    machine_id = data.get("machine_id")
    if not machine_id:
        return jsonify({"error": "Missing machine_id"}), 400

    machine = get_machine(machine_id)
    if not machine:
        return jsonify({"error": "Machine not found"}), 404

    if "esp32_ip" in data:
        candidate_ip = str(data.get("esp32_ip") or "").strip()
        if not candidate_ip:
            return jsonify({"error": "ESP32 IP cannot be empty."}), 400
        if not _is_valid_ipv4(candidate_ip):
            return jsonify({"error": "ESP32 IP must be a valid IPv4 address."}), 400

    try:
        updated = update_machine(
            machine_id=machine_id,
            name=(str(data.get("name") or "").strip() if "name" in data else None),
            esp32_ip=(str(data.get("esp32_ip") or "").strip() if "esp32_ip" in data else None),
            machine_type=(str(data.get("type") or "").strip() if "type" in data else None),
            machine_function=(str(data.get("machine_function") or "").strip() if "machine_function" in data else None),
            pulse_on=(int(data["pulse_on"]) if "pulse_on" in data else None),
            pulse_off=(int(data["pulse_off"]) if "pulse_off" in data else None),
            pulse_count=(int(data["pulse_count"]) if "pulse_count" in data else None),
            vend_price=(int(data["vend_price"]) if "vend_price" in data else None),
            quick_wash_pulse_count=(int(data["quick_wash_pulse_count"]) if "quick_wash_pulse_count" in data else None),
            quick_wash_price=(int(data["quick_wash_price"]) if "quick_wash_price" in data else None),
            custom_modes=(data.get("custom_modes") if "custom_modes" in data else None),
            order_index=(int(data["order_index"]) if "order_index" in data else None),
        )
    except (TypeError, ValueError):
        return jsonify({"error": "Invalid numeric settings."}), 400

    if not updated:
        return jsonify({"error": "No settings provided."}), 400

    return jsonify({"status": "ok"}), 200


@dashboard_bp.route("/dashboard/machines/delete", methods=["POST"])
def delete_machines_route():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json() or {}
    machine_ids = data.get("machine_ids", [])
    
    if not machine_ids or not isinstance(machine_ids, list):
        return jsonify({"error": "Missing or invalid machine_ids array"}), 400

    deleted = delete_machines(machine_ids)
    return jsonify({"status": "ok", "deleted": deleted}), 200


@dashboard_bp.route("/dashboard/machines/live-status")
def live_machine_status():
    machines = get_all_machines()
    status_map = {}
    for m in machines:
        status_map[m["id"]] = get_esp32_status(m["esp32_ip"])
    return jsonify(status_map)


@dashboard_bp.route("/dashboard/machines/runtime")
def machine_runtime():
    machines = get_all_machines()
    runtime_map = {}
    for m in machines:
        normalized = _normalize_machine_runtime(m)
        runtime_map[normalized["id"]] = {
            "status": normalized.get("status"),
            "run_started_at": normalized.get("run_started_at"),
            "run_ends_at": normalized.get("run_ends_at"),
            "remaining_seconds": int(normalized.get("remaining_seconds") or 0),
        }
    return jsonify(runtime_map)


@dashboard_bp.route("/dashboard/transactions/delete", methods=["POST"])
def dashboard_delete_transactions():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json() or {}
    transaction_ids = data.get("transaction_ids") or []
    if not isinstance(transaction_ids, list):
        return jsonify({"error": "transaction_ids must be a list"}), 400

    result = delete_transactions(transaction_ids)
    return jsonify(
        {
            "status": "ok",
            "deleted_count": result.get("deleted_count", 0),
            "restocked_quantity": result.get("restocked_quantity", 0),
        }
    )


@dashboard_bp.route("/dashboard/transactions/latest", methods=["GET"])
def dashboard_get_latest_transaction():
    txns = _enrich_transactions(_get_transactions(limit=200))
    latest = None
    for txn in txns:
        if str(txn.get("status") or "").upper() not in ("COMPLETED", "SIMULATED"):
            continue
        latest = txn
        break

    if not latest:
        return jsonify({"error": "No completed transactions found"}), 404

    return jsonify(
        {
            "transaction_id": latest.get("id"),
            "amount": int(latest.get("amount") or 0),
        }
    )


@dashboard_bp.route("/dashboard/api/transactions", methods=["GET"])
def get_transactions_api():
    limit = request.args.get("limit", default=2000, type=int)
    offset = request.args.get("offset", default=0, type=int)
    transactions = _enrich_transactions(_get_transactions(limit=limit, offset=offset))
    return jsonify(transactions)


@dashboard_bp.route("/dashboard/transactions/find-by-amount", methods=["GET"])
def dashboard_find_transaction_by_amount():
    amount_arg = request.args.get("amount")
    try:
        target_amount = int(amount_arg)
    except (TypeError, ValueError):
        return jsonify({"error": "amount query param must be a whole number"}), 400

    if target_amount < 0:
        return jsonify({"error": "amount must be 0 or higher"}), 400

    txns = _enrich_transactions(_get_transactions(limit=500))
    matched = []
    running_total = 0
    for txn in txns:
        if str(txn.get("status") or "").upper() not in ("COMPLETED", "SIMULATED"):
            continue
        amount = int(txn.get("amount") or 0)
        if txn.get("gcash_amount") is not None:
            current_gcash = int(txn.get("gcash_amount") or 0)
        else:
            current_gcash = amount if int(txn.get("paid_by_gcash") or 0) == 1 else 0
        current_gcash = max(0, min(current_gcash, amount))
        cash_remaining = amount - current_gcash
        if cash_remaining <= 0:
            continue
        if running_total + cash_remaining > target_amount:
            continue

        matched.append(
            {
                "transaction_id": txn.get("id"),
                "amount": amount,
                "cash_remaining": cash_remaining,
            }
        )
        running_total += cash_remaining
        if running_total == target_amount:
            break

    if running_total != target_amount or not matched:
        return jsonify({"error": "No combination of recent transactions matched this amount"}), 404

    return jsonify(
        {
            "matched_total": running_total,
            "transaction_count": len(matched),
            "transactions": matched,
        }
    )


@dashboard_bp.route("/dashboard/transactions/<transaction_id>/payment-method", methods=["POST"])
def dashboard_update_transaction_payment_method(transaction_id):
    data = request.get_json(silent=True) or {}
    if "gcash_amount" in data:
        requested_gcash = data.get("gcash_amount")
    elif "paid_by_gcash" in data:
        requested_gcash = None if not bool(data.get("paid_by_gcash")) else -1
    else:
        return jsonify({"error": "Missing gcash_amount"}), 400

    if requested_gcash == -1:
        conn = get_connection()
        row = conn.execute("SELECT amount FROM transactions WHERE id = ? LIMIT 1", (transaction_id,)).fetchone()
        conn.close()
        if not row:
            return jsonify({"error": "Transaction not found"}), 404
        requested_gcash = int(row["amount"] or 0)

    try:
        updated = update_transaction_gcash_amount(transaction_id, requested_gcash)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    if not updated:
        return jsonify({"error": "Transaction not found"}), 404

    return jsonify(
        {
            "status": "ok",
            "transaction_id": updated.get("id"),
            "amount": int(updated.get("amount") or 0),
            "paid_by_gcash": int(updated.get("paid_by_gcash") or 0),
            "gcash_amount": int(updated.get("gcash_amount") or 0),
            "cash_amount": int(updated.get("cash_amount") or 0),
        }
    )


@dashboard_bp.route("/dashboard/export/sheets")
def export_sheets():
    import openpyxl
    from openpyxl.styles import Font, Alignment
    
    machines = _enrich_machines(get_all_machines())
    transactions = _enrich_transactions(_get_transactions())

    # Optional filters
    status_filter = request.args.get("status")
    search_filter = request.args.get("search", "").lower()
    sort_dir = request.args.get("sort", "desc")
    
    start_str = request.args.get("start")
    end_str = request.args.get("end")
    start_date = datetime.strptime(start_str, "%Y-%m-%d").date() if start_str else None
    end_date = datetime.strptime(end_str, "%Y-%m-%d").date() if end_str else None

    # Filter by date range first
    transactions = _filter_by_date(transactions, start_date, end_date)

    if status_filter:
        transactions = [t for t in transactions if t.get("status") == status_filter]
    if search_filter:
        transactions = [
            t for t in transactions
            if search_filter in t.get("id", "").lower()
            or search_filter in t.get("machine_id", "").lower()
            or search_filter in t.get("location_name", "").lower()
        ]

    # Sort by date
    transactions.sort(
        key=lambda t: t.get("started_at", ""),
        reverse=True if sort_dir == "desc" else False,
    )

    machine_name_by_id = {m["id"]: m.get("name", m["id"]) for m in machines}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Write header
    headers = ["Transaction ID", "Timestamp", "Machine", "Location", "Amount", "Status"]
    ws.append(headers)
    
    # Style header
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Set column widths
    ws.column_dimensions['A'].width = 38  # Transaction ID
    ws.column_dimensions['B'].width = 22  # Timestamp
    ws.column_dimensions['C'].width = 15  # Machine
    ws.column_dimensions['D'].width = 15  # Location
    ws.column_dimensions['E'].width = 10  # Amount
    ws.column_dimensions['F'].width = 15  # Status

    # Data rows
    for t in transactions:
        ws.append([
            t.get("id", ""),
            t.get("started_at", ""),
            machine_name_by_id.get(t.get("machine_id", ""), t.get("machine_id", "")),
            t.get("location_name", ""),
            t.get("amount", 0),
            t.get("status", "")
        ])

    byte_buffer = io.BytesIO()
    wb.save(byte_buffer)
    byte_buffer.seek(0)
    
    dates = []
    for t in transactions:
        try:
            dates.append(datetime.strptime(t["started_at"], "%Y-%m-%d %H:%M:%S").date())
        except (ValueError, TypeError):
            pass
    earliest = min(dates) if dates else datetime.now().date()
    latest = max(dates) if dates else datetime.now().date()

    filename = f"LaundryLink_Sales_{earliest}_{latest}.xlsx"

    return send_file(
        byte_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Email Settings Routes
# ---------------------------------------------------------------------------

@dashboard_bp.route("/dashboard/settings/email", methods=["GET"])
def get_email_settings_route():
    guard = require_admin_pin()
    if guard:
        return guard

    settings = get_email_settings()
    return jsonify(settings)


@dashboard_bp.route("/dashboard/settings/email", methods=["POST"])
def update_email_settings_route():
    guard = require_admin_pin()
    if guard:
        return guard

    data = request.get_json(silent=True) or {}



    try:
        updated = set_email_settings(data)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    # Reschedule the APScheduler job to match the new settings
    try:
        from services.emailing import refresh_email_schedule
        refresh_email_schedule()
    except Exception:
        pass  # Scheduler not running in test environments

    return jsonify({"status": "ok", **updated}), 200


@dashboard_bp.route("/dashboard/settings/email/test", methods=["POST"])
def trigger_test_email_route():
    guard = require_admin_pin()
    if guard:
        return guard

    settings = get_email_settings()
    if not settings.get("smtp_email") or not settings.get("smtp_password"):
        return jsonify({"error": "SMTP credentials are not configured."}), 400
    if not settings.get("recipients"):
        return jsonify({"error": "No recipients configured."}), 400

    try:
        from services.emailing import trigger_test_email
        trigger_test_email()
    except Exception as exc:
        return jsonify({"error": f"Failed to trigger test email: {exc}"}), 500

    return jsonify({"status": "ok", "message": "Test email is being sent in the background."}), 200
